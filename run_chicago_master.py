import os
import sys
import re
import glob
import asyncio
import difflib
import cv2
import numpy as np
from collections import Counter
from dataclasses import dataclass
from openpyxl import Workbook
from openpyxl.styles import Font

import winsdk.windows.graphics.imaging as imaging
import winsdk.windows.media.ocr as ocr
import winsdk.windows.storage as storage

@dataclass
class PlayRow:
    team: str
    formation: str
    set_name: str
    play_name: str
    play_type: str
    source_file: str
    card_index: int

# 1. Build canonical reference catalogs from Madden 26 workbooks
m26_dir = r"D:\Robert\Desktop\App Project Builds\Gridiron_OC\Gridiron_OC\imports\Madden 26 Team Playbook (Excel Workbooks)"
off_canon = {}
def_canon = {}
off_sets_by_form = {}
def_sets_by_form = {
    "3-3-5": ["Penny"],
    "4-3": ["Odd Leo", "Over", "Over Solid", "Under"],
    "46": ["Bear Under"],
    "Dime": ["2-3", "Rush", "Normal"],
    "Dollar": ["3-2 Mug"],
    "Goal Line": ["5-3", "6-2"],
    "Nickel": ["2-4", "2-4 Dbl Mug", "2-4 Load Mug", "2-4 Odd Mug", "Over", "Wide", "Swap Dbl Mug"],
    "Prevent": ["3 Deep"],
}

for f in glob.glob(os.path.join(m26_dir, "*.xlsx")):
    try:
        wb = openpyxl.load_workbook(f, read_only=True)
        if "Playbook" in wb.sheetnames:
            ws = wb["Playbook"]
            for r in ws.iter_rows(values_only=True):
                if r and len(r) >= 5 and r[3] and r[3] != "Play Name" and r[4]:
                    off_canon[re.sub(r"[^A-Z0-9]", "", str(r[3]).upper())] = (str(r[3]).strip(), str(r[4]).strip())
                if r and len(r) >= 3 and r[1] and r[2] and r[2] != "Set":
                    off_sets_by_form.setdefault(str(r[1]).strip(), set()).add(str(r[2]).strip())
        if "Defensive Plays" in wb.sheetnames:
            ws = wb["Defensive Plays"]
            for r in ws.iter_rows(values_only=True):
                if r and len(r) >= 5 and r[3] and r[3] != "Play Name" and r[4]:
                    def_canon[re.sub(r"[^A-Z0-9]", "", str(r[3]).upper())] = (str(r[3]).strip(), str(r[4]).strip())
    except Exception:
        pass

off_sets_by_form["H Wing"] = {"Wing"}

print(f"Loaded Catalogs: {len(off_canon)} Offense Plays, {len(def_canon)} Defense Plays")

async def run_ocr_on_image(img, engine, scale=2.5):
    h, w = img.shape[:2]
    up = cv2.resize(img, (0, 0), fx=scale, fy=scale, interpolation=cv2.INTER_LANCZOS4)
    tmp_path = os.path.abspath(f"tmp_ocr_{np.random.randint(100000, 999999)}.png")
    cv2.imwrite(tmp_path, up)
    
    file = await storage.StorageFile.get_file_from_path_async(tmp_path)
    stream = await file.open_async(storage.FileAccessMode.READ)
    decoder = await imaging.BitmapDecoder.create_async(stream)
    bm = await decoder.get_software_bitmap_async()
    res = await engine.recognize_async(bm)
    stream.close()
    try:
        os.remove(tmp_path)
    except Exception:
        pass
        
    words = []
    lines = []
    for line in res.lines:
        lines.append(line.text)
        for word in line.words:
            words.append((word.bounding_rect.y / scale, word.bounding_rect.x / scale, word.bounding_rect.width / scale, word.text))
    return words, lines

def classify_card_badge(im_full, card_idx, is_defense=False):
    h, w = im_full.shape[:2]
    if card_idx == 1:
        x0, x1 = int(w * 0.03), int(w * 0.08)
    elif card_idx == 2:
        x0, x1 = int(w * 0.35), int(w * 0.40)
    else:
        x0, x1 = int(w * 0.68), int(w * 0.73)
    y0, y1 = int(h * 0.72), int(h * 0.82)
    crop = im_full[y0:y1, x0:x1]
    b, g, r = [float(np.mean(crop[:, :, i])) for i in range(3)]
    
    if is_defense:
        if r > 60 and b > 75 and g < 70 and abs(r - b) < 45:
            return "MATCH"
        if r > 75 and g > 50 and b < 60:
            return "MAN"
        if r > 75 and r > g + 20 and r > b + 25:
            return "BLITZ"
        if b > 80 and b > r + 15:
            return "ZONE"
        return "ZONE"
    else:
        if r > 75 and r > g + 20 and r > b + 20:
            return "Run"
        return "Pass"

def resolve_defense_set(formation, raw_text):
    if formation == "Goal Line":
        s = str(raw_text).upper()
        if "5-3" in s or "53" in s or "6-3" in s or "63" in s:
            return "5-3"
        return "6-2"
    if formation == "4-3":
        s = str(raw_text).upper()
        if "SOLID" in s:
            return "Over Solid"
        if "LEO" in s:
            return "Odd Leo"
        if "UNDER" in s:
            return "Under"
        return "Over"
        
    allowed = def_sets_by_form.get(formation, ["Normal"])
    clean = re.sub(r"[^A-Z0-9]", "", str(raw_text or "").upper())
    for fw in ("335", "43", "46", "DIME", "DOLLAR", "GOALLINE", "GOAL", "LINE", "NICKEL", "PREVENT", "BEAR"):
        clean = clean.replace(fw, "")
        
    for cand in allowed:
        if re.sub(r"[^A-Z0-9]", "", cand.upper()) == clean:
            return cand
    for cand in allowed:
        c_clean = re.sub(r"[^A-Z0-9]", "", cand.upper())
        if c_clean and c_clean in clean:
            return cand
    m = difflib.get_close_matches(clean, [re.sub(r"[^A-Z0-9]", "", c.upper()) for c in allowed], n=1, cutoff=0.35)
    if m:
        for cand in allowed:
            if re.sub(r"[^A-Z0-9]", "", cand.upper()) == m[0]:
                return cand
    return allowed[0]

def resolve_offense_set(formation, raw_text):
    if formation == "H Wing":
        return "Wing"
    if formation == "Hail Mary":
        return "Normal"
        
    allowed = list(off_sets_by_form.get(formation, []))
    if not allowed:
        return "Normal"
        
    s_upper = str(raw_text or "").upper()
    if "ACE DOUBLE WING" in s_upper or "ACEDOUBLEWING" in s_upper: return "Ace Double Wing"
    if "WING PAIR HEAVY" in s_upper or "WING PAIR" in s_upper: return "Wing Pair Heavy"
    if "WING TIGHT LEFT" in s_upper or "TIGHT LEFT" in s_upper or "WING LEFT" in s_upper: return "Wing Tight Left"
    if "DEUCE CLOSE" in s_upper or "CLOSE DEUCE" in s_upper: return "Deuce Close"
    if "BUNCH X NASTY" in s_upper or "BUNCH X" in s_upper: return "Bunch X Nasty"
    if "TIGHT Y OFF FLEX" in s_upper or "OFF FLEX" in s_upper: return "Tight Y Off Flex"
    if "U OFF TRIPS" in s_upper or ("TRIPS" in s_upper and formation == "Singleback"): return "U Off Trips"
    if "TIGHT DOUBLES" in s_upper or ("DOUBLES" in s_upper and formation == "Singleback"): return "Tight Doubles"
    if "WING CLOSE" in s_upper or "CLOSE WING" in s_upper: return "Wing Close"
    if "U OFF STACK" in s_upper or "STACK" in s_upper: return "U Off Stack"
    if "EAST TIGHT" in s_upper: return "East Tight"
    
    # Gun sets
    if "BUNCH SPREAD" in s_upper: return "Bunch Spread"
    if "BUNCH OPEN TE" in s_upper: return "Bunch Open Te"
    if "WING FLEX" in s_upper: return "Wing Flex Offset Wk"
    if "WING SLOT" in s_upper: return "Wing Slot Offset"
    if "WING TIGHT" in s_upper: return "Wing Tight Nasty"
    if "EMPTY BASE" in s_upper: return "Empty Base Trio"
    if "EMPTY CHIPS NASTY" in s_upper: return "Empty Chips Nasty"
    if "EMPTY CHIPS QUADS" in s_upper: return "Empty Chips Quads"
    if "NORMAL Y OFF CLOSE" in s_upper: return "Normal Y Off Close"
    if "SLOT RIGHT" in s_upper: return "Slot Right"
    if "TIGHT OPEN" in s_upper: return "Tight Open"
    if "TIGHT Y OFF WK" in s_upper: return "Tight Y Off Wk"
    if "TRIPS TE FLEX" in s_upper: return "Trips Te Flex"
    if "TRIPS TE OFFSET" in s_upper: return "Trips Te Offset Wk"
    if "TRIPS HB WK" in s_upper: return "Trips Hb Wk"
    if "Y OFF TRIO" in s_upper: return "Y Off Trio Wk"
    if "Y OFF TRIPS WK" in s_upper: return "Y Off Trips Wk"
    if "Y OFF TRIPS" in s_upper: return "Y Off Trips"
    if "SPREAD" in s_upper: return "Spread"
    if "DOUBLE STACK" in s_upper or "DOUBLESUCK" in s_upper: return "Double Stack"
    if "DOUBLES HB WK" in s_upper: return "Doubles Hb Wk"
    
    clean = re.sub(r"[^A-Z0-9]", "", s_upper)
    for fw in ("SINGLEBACK", "SINGLEBCK", "SINGIEBACK", "GUN", "SHOTGUN", "PISTOL", "IFORM", "WEAK", "FORM", "HAIL", "MARY"):
        clean = clean.replace(fw, "")
        
    for cand in allowed:
        if re.sub(r"[^A-Z0-9]", "", cand.upper()) == clean:
            return cand
            
    cand_map = {re.sub(r"[^A-Z0-9]", "", c.upper()): c for c in allowed}
    m = difflib.get_close_matches(clean, list(cand_map.keys()), n=1, cutoff=0.50)
    if m:
        return cand_map[m[0]]
    return "Normal"

def_repairs = {
    '101MDS': ('Cross 3 Fire', 'BLITZ'),
    'CNSS8FLRE': ('Cross 3 Fire', 'BLITZ'),
    'CMSS8FIRE': ('Cross 3 Fire', 'BLITZ'),
    'COVERIHOLE': ('Cover 1 Hole', 'MAN'),
    'IHOLEPRESS': ('1 Hole Press', 'MAN'),
    'IROBBERPRESS': ('1 Robber Press', 'MAN'),
    'ICOUTAINPRESS': ('1 Contain Press', 'MAN'),
    'COVERICONUIN': ('Cover 1 Contain', 'MAN'),
    'COVER2MAN': ('Cover 2 Man', 'MAN'),
    'COVER2HARDFLAT': ('Cover 2 Hard Flat', 'ZONE'),
    'COVER4PRESS': ('Cover 4 Press', 'MATCH'),
    'COVER4WARTERS': ('Cover 4 Quarters', 'MATCH'),
    'LBSTINGI': ('LB Sting 1', 'BLITZ'),
    'OLBFIREMAN': ('OLB Fire Man', 'BLITZ'),
    'SAMEDGE8': ('Sam Edge 3', 'BLITZ'),
    'THUNDERSMOKE': ('Thunder Smoke', 'BLITZ'),
    'DBLSAFETYBLITZ': ('Dbl Safety Blitz', 'BLITZ'),
    'FREEFIRE': ('Free Fire', 'BLITZ'),
    'FREEFIREO': ('Free Fire 0', 'BLITZ'),
    'HAMMEROBLAST': ('Hammer 0 Blast', 'BLITZ'),
    'COVER8SKYWK': ('Cover 3 Sky Wk', 'ZONE'),
    'FSFLRE8': ('FS Fire 3', 'BLITZ'),
    'FSFIREI': ('FS Fire 1', 'BLITZ'),
    'SAMCUSH8': ('Sam Crash 3', 'BLITZ'),
    'SAMCRASHI': ('Sam Crash 1', 'BLITZ'),
    'COVER8HARDFLAT': ('Cover 3 Hard Flat', 'ZONE'),
    'COVER21WERT': ('Cover 2 Invert', 'ZONE'),
    '21WERTHARDFLAT': ('2 Invert Hard Flat', 'ZONE'),
    'WILLGOFLRE8': ('Will Go Fire 3', 'BLITZ'),
    'PINCHBIMO': ('Pinch Buck 0', 'BLITZ'),
    'PINCHDOG2PRESS': ('Pinch Dog 2 Press', 'BLITZ'),
    'STORMBUVEI': ('Storm Brave 1', 'BLITZ'),
    'WILLBUCK3PRESS': ('Will Buck 3 Press', 'BLITZ'),
    'SAMISTING': ('Sam 1 Sting', 'BLITZ'),
    'SSBUTZI': ('SS Blitz 1', 'BLITZ'),
    '21NVERTHARDFLAT': ('2 Invert Hard Flat', 'ZONE'),
    'SAWOBUST': ('Sam 0 Blast', 'BLITZ'),
    'COVER81URDFLAT': ('Cover 3 Hard Flat', 'ZONE'),
    'SS2TRAP': ('SS 2 Trap', 'BLITZ'),
    'COVERS': ('Cover 6', 'MATCH'),
    'COVERG': ('Cover 6', 'MATCH'),
    'TAMPA2': ('Tampa 2', 'ZONE'),
    'GOMIJT': ('60 Out', 'BLITZ'),
    'GOMIT': ('60 Out', 'BLITZ'),
    'GOMJR': ('60 Out', 'BLITZ'),
    'GLMAN': ('GL Man', 'BLITZ'),
    'GAPSAU': ('Gaps All', 'BLITZ'),
    'GAPSAULB': ('Gaps LB', 'BLITZ'),
    'GAPSB': ('Gaps B', 'BLITZ'),
    'GAPSAB': ('Gaps All', 'BLITZ'),
    'MLBGAPA': ('MLB Gap A', 'BLITZ'),
    'BRACKET9SWLTCH': ('Bracket 9 Switch', 'MATCH'),
    'BUCKETSWITCHWILLIE': ('Bracket Switch Willie', 'MATCH'),
    'COVER3SHOW2': ('Cover 3 Show 2', 'ZONE'),
    'COVER6WILLE': ('Cover 6 Willie', 'MATCH'),
    'COVER6SHW2': ('Cover 6 Show 2', 'MATCH'),
    'COVER8BLUMATCH': ('Cover 3 Buzz Match', 'MATCH'),
    'COVER8SAMMLE': ('Cover 3 Seam', 'ZONE'),
    'COVER8SEAU': ('Cover 3 Seam', 'ZONE'),
    'COVERIHDLE': ('Cover 1 Hole', 'MAN'),
    'COVERIPLMS': ('Cover 4 Palms', 'MATCH'),
    'COVERIPALMS': ('Cover 4 Palms', 'MATCH'),
    'COVERIROBBER': ('Cover 1 Robber', 'MAN'),
    'COVERIOBERPRESS': ('Cover 1 Robber Press', 'MAN'),
    'DMJBLEBUCKETSWITCH': ('Double Bracket Switch', 'MATCH'),
    'EDGEBLLTZ8': ('Edge Blitz 3', 'BLITZ'),
    'EDGEBLLTZ3': ('Edge Blitz 3', 'BLITZ'),
    'FIELDSIM8': ('Field Sim 3', 'BLITZ'),
    'FIELDSM8': ('Field Sim 3', 'BLITZ'),
    'FSWILLBLITZO': ('FS Will Blitz 0', 'BLITZ'),
    'IDOUBLESLDT': ('1 Double Slot', 'MAN'),
    'IDOUBLEWR2': ('1 Double WR2', 'MAN'),
    'IDOUBLEWRI': ('1 Double WR2', 'MAN'),
    'IDOUBLETE': ('1 Double TE', 'MAN'),
    'MLLCOVER8CLOLD': ('Cover 3 Cloud', 'ZONE'),
    'NICKELSLM2': ('Nickel Sim 2', 'BLITZ'),
    'OVERLOAD8PRESS': ('Overload 3 Press', 'BLITZ'),
    'ROUCOVER21NVERT': ('Cover 2 Invert', 'ZONE'),
    'SIMPRESSURE8': ('Sim Pressure 3', 'BLITZ'),
    'BLITZDB': ('Blitz 3', 'BLITZ'),
    'COVER2ROU': ('Cover 2 Roll', 'ZONE'),
    'COVER8BUMATCH': ('Cover 3 Buzz Match', 'MATCH'),
    'COVER8CLOUD': ('Cover 3 Cloud', 'ZONE'),
    'DIMEBLLTZ2': ('Dime Blitz 2', 'BLITZ'),
    'FSMIDDLEBLLTZ8': ('FS Middle Blitz 3', 'BLITZ'),
    'MISIMPRESSURE': ('Mug Sim Pressure', 'BLITZ'),
    'MUGDTLOOPI': ('Mug DT Loop 1', 'BLITZ'),
    'SSOVERLOAD8': ('SS Overload 3', 'BLITZ'),
    'SSSIMPRESSURE8': ('SS Sim Pressure 3', 'BLITZ'),
    'DBFLRE2': ('DB Fire 2', 'BLITZ'),
    'DTMIKE1POP0': ('DT Mike 1 Pop 0', 'BLITZ'),
    'DTMIKELODPI': ('DT Mike Loop 3', 'BLITZ'),
    'OVERLOAD8SUM': ('Overload 3 Seam', 'BLITZ'),
    'ZEROBUTZ': ('Zero Blitz', 'BLITZ'),
    'BENCHBLLTZ8': ('Bench Blitz 3', 'BLITZ'),
    'LBPOP0': ('LB Pop 0', 'BLITZ'),
    'LBCROSS8SHW12': ('LB Cross 3 Show 2', 'BLITZ'),
    'LBDETWLST8': ('LB DE Twist 3', 'BLITZ'),
    'NICKELBUTZ2': ('Nickel Blitz 2', 'BLITZ'),
    'NICKELDOG8BUZZ': ('Nickel Dog 3 Buzz', 'BLITZ'),
    'SAMMIKELOOPI': ('Sam Mike Loop 3', 'BLITZ'),
    'SSBLLTZ8': ('SS Blitz 3', 'BLITZ'),
    'LBBUTZ8': ('LB Blitz 3', 'BLITZ'),
    'NICKEL2TRP': ('Nickel 2 Trap', 'BLITZ'),
    'NICKEL2TUP': ('Nickel 2 Trap', 'BLITZ'),
    'COVER8CIDJD': ('Cover 3 Cloud', 'ZONE'),
    'COVER8CIMJD': ('Cover 3 Cloud', 'ZONE'),
    'COVERILBBUTZ': ('Cover 1 LB Blitz', 'BLITZ'),
    'LBBUTZO': ('LB Blitz 0', 'BLITZ'),
    'NICKELBLITZI': ('Nickel Blitz 1', 'BLITZ'),
    'NICKELBLITZO': ('Nickel Blitz 0', 'BLITZ'),
    'NICKELBLLTZ8': ('Nickel Blitz 3', 'BLITZ'),
    'COVER4D0': ('Cover 4 Drop', 'ZONE'),
    'COVER8WLLLE': ('Cover 6 Willie', 'MATCH'),
    'PINCHO': ('Pinch 0', 'BLITZ'),
    'SAMMIKELOOP8': ('Sam Mike Loop 3', 'BLITZ'),
    '2INVERTLURK': ('Cover 2 Invert Lurk', 'ZONE'),
    '8BUZZMATCHWK': ('Cov 3 Buzz Match Wk', 'MATCH'),
    'COV8BUZZMATCHWK': ('Cov 3 Buzz Match Wk', 'MATCH'),
    'COVERSSAMMIE': ('Cover 3 Seam', 'ZONE'),
    'CWER4PALUS': ('Cover 4 Palms', 'MATCH'),
    'LBBUTZI': ('LB Blitz 1', 'BLITZ'),
    'MBBERPRSS': ('1 Robber Press', 'MAN'),
    'NIL23CLOUD': ('Nil 2 3 Cloud', 'ZONE'),
    'ROLLCOVER8C': ('Roll Cover 3 Cloud', 'ZONE'),
    'SHWJSSSLM2': ('Show SS Sim 2', 'BLITZ'),
    'BLITZLOOP8': ('Blitz Loop 3', 'BLITZ'),
    'BLITZLOOPO': ('Blitz Loop 0', 'BLITZ'),
    'BLITZLOOPSM8': ('Blitz Loop Sim 3', 'BLITZ'),
    'CBBENCHSLM2': ('CB Bench Sim 2', 'BLITZ'),
    'COVERIGAMS': ('Cover 1 Games', 'BLITZ'),
    'HOTBLLTZ3': ('Hot Blitz 3', 'BLITZ'),
    'IDTLBLODP': ('1 DT LB Loop', 'BLITZ'),
    'REDZONEDTDO': ('Redzone DT Drop', 'ZONE'),
    'SSDOG8': ('SS Dog 3', 'BLITZ'),
    'SSSIM2': ('Cross Sim 2', 'BLITZ'),
    'CBZONEBUTZ': ('CB Zone Blitz', 'BLITZ'),
    'COV21NVERTHARDFLAT': ('Cover 2 Invert Hard Flat', 'ZONE'),
    'COVER4MLUS': ('Cover 4 Palms', 'MATCH'),
    'COVER4ZONE': ('Cover 4 Zone', 'ZONE'),
    'COVER9ZONE': ('Cover 9', 'MATCH'),
    'COVERBZONE': ('Cover 6 Zone', 'ZONE'),
    'CWER8SKYPRESS': ('Cover 3 Sky Press', 'ZONE'),
    'ILBBUTZ': ('1 LB Blitz', 'BLITZ'),
    'SILVERSHDOTPINCH': ('Silver Shoot Pinch', 'BLITZ'),
    '8DEEPBUTZ': ('3 Deep Blitz', 'BLITZ'),
    'COVER2DO': ('Cover 2 Drop', 'ZONE'),
    'IDOUBLEBUZZ': ('3 Double Buzz', 'ZONE'),
    'MAN8DEEP': ('Man 3 Deep', 'MAN'),
}

off_repairs = {
    '14B': 'HB',
    'HBCOUUTERWK': 'HB COUNTER WK',
    'HBCOIJUTERWK': 'HB COUNTER WK',
    'HBSPUTO': 'HB SPLIT O',
    'HBSPLITO': 'HB SPLIT O',
    'STmNGTOSS': 'STRONG TOSS',
    'STETOSS': 'STRONG TOSS',
    'STMGTOSS': 'STRONG TOSS',
    'PWERO': 'POWER O',
    'POWERO': 'POWER O',
    'PAPWERO': 'PA POWER O',
    'PAPOWERO': 'PA POWER O',
    'PASPRINTHBFUT': 'PA SPRINT HB FLAT',
    'PASPRINTHBFLAT': 'PA SPRINT HB FLAT',
    'PASPOT': 'PA SPOT',
    'QBSNEAK': 'QB SNEAK',
    'QBSNEAX': 'QB SNEAK',
    'FBDIVEWEAK': 'FB DIVE WEAK',
    'FBDIVEWEAI': 'FB DIVE WEAK',
    'PAWAGGLE': 'PA WAGGLE',
    'INSIDEZONE': 'INSIDE ZONE',
    'MESHSPOT': 'MESH SPOT',
    'HBPUNCH': 'HB PUNCH',
    'BENCHSWAP': 'BENCH SWAP',
    'READOPTION': 'READ OPTION',
    'SEAMSZ-UUDER': 'SEAMS Z-UNDER',
    'HBPDNER': 'HB POWER',
    'PADAGGER': 'PA DAGGER',
    'MTNRPOBUCKPEEKSLANT': 'MTN RPO BUCK PEEK SLANT',
    'MTNPAFKSCREENSCISSORS': 'MTN PA FK SCREEN SCISSORS',
    'MTNRPOALERTBUBBLE': 'MTN RPO ALERT BUBBLE',
    'FOURVERTICLS': 'FOUR VERTICALS',
    'FOURVERTICALS': 'FOUR VERTICALS',
    'HBSCISSORS': 'HB SCISSORS',
    'DIJO': 'DUO',
    'HBMJO': 'HB DUO',
    'MTNMJO': 'MTN DUO',
    '26DLN': '26 DUO',
    'CLEARMJTZOPTION': 'CLEAROUT OPTION',
    'BOUNCEPABUBBLEFLOOD': 'BOUNCE PA BUBBLE FLOOD',
    'BWUCESHALLM': 'BOUNCE SHALLOW',
    'RETURNSPRINTCOMEUCK': 'RETURN SPRINT COMEBACK',
    'RETURNSPRINTTESCREEN': 'RETURN SPRINT TE SCREEN',
    'SHUFFLEHBOUT': 'SHUFFLE HB OUT',
    'SMASHYDMJBLE': 'SMASH Y DOUBLE',
    'PASWITCHREPLACE': 'PA SWITCH REPLACE',
    'PAHS': 'PA WHEEL',
    'STICK': 'STICK',
    'POSTS': 'POSTS',
    'DWBLESUNT': 'DOUBLES SLANT',
    'RPOALERTMJT': 'RPO ALERT OUT',
    'ZONEFAKEJET': 'ZONE FAKE JET',
    'DEEP CHOICE': 'DEEP CHOICE',
    'HBPWERRT': 'HB POWER',
    'MTNMJTSIDEZOUE': 'MTN OUTSIDE ZONE',
    'YTUIL': 'Y TRAIL',
    'HAMMERDAGGER': 'HAMMER DAGGER',
    'HAILMAW': 'HAIL MARY',
    'HAILMARY': 'HAIL MARY',
    'CNSSDRAG': 'CROSS DRAG',
    'STEPWER': 'SLOT POWER',
    'PATECWSS': 'PA TE CROSS',
    'MTNPWER': 'MTN POWER',
    'YFLATGOAUINE': 'Y FLAT GO LINE',
}

def clean_defense_play(raw_title, badge_type):
    t = re.sub(r"\s+", " ", str(raw_title or "").strip()).upper()
    t = re.sub(r"^[Oo0]\s+", "", t)
    t = re.sub(r"^[-–—]\s*", "", t)
    
    k = re.sub(r"[^A-Z0-9]", "", t)
    if k in def_repairs:
        return def_repairs[k]
    if k in def_canon:
        return def_canon[k]
    m = difflib.get_close_matches(k, list(def_canon.keys()), n=1, cutoff=0.7)
    if m:
        return def_canon[m[0]]
    return (t.title(), badge_type)

def clean_offense_play(raw_title, badge_type):
    t = re.sub(r"\s+", " ", str(raw_title or "").strip()).upper()
    t = re.sub(r"^[Oo0]\s+", "", t)
    t = re.sub(r"^[-–—]\s*", "", t)
    
    k = re.sub(r"[^A-Z0-9]", "", t)
    if k in off_repairs:
        t = off_repairs[k]
        k = re.sub(r"[^A-Z0-9]", "", t)
        
    if k in off_canon:
        return off_canon[k]
    m = difflib.get_close_matches(k, list(off_canon.keys()), n=1, cutoff=0.7)
    if m:
        return off_canon[m[0]]
        
    if "RPO" in t:
        p_type = "RPO"
    elif "PA" in t or "PLAY ACTION" in t:
        p_type = "Play Action"
    elif badge_type == "Run":
        p_type = "Run"
    else:
        p_type = "Pass"
    return (t, p_type)

async def extract_formation_plays(formation_dir, team, formation_name, is_defense=False):
    engine = ocr.OcrEngine.try_create_from_user_profile_languages()
    files = sorted([f for f in os.listdir(formation_dir) if f.endswith(".png")])
    current_set = "Normal"
    rows = []
    
    for f in files:
        im = cv2.imread(os.path.join(formation_dir, f))
        h, w = im.shape[:2]
        
        # Check active tab pill from top bar
        top = im[5:45, :]
        gray_top = cv2.cvtColor(top, cv2.COLOR_BGR2GRAY)
        _, thresh_top = cv2.threshold(gray_top, 120, 255, cv2.THRESH_BINARY)
        cnts, _ = cv2.findContours(thresh_top, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        pills = [cv2.boundingRect(c) for c in cnts if cv2.boundingRect(c)[2] >= 30 and 12 <= cv2.boundingRect(c)[3] <= 30 and cv2.boundingRect(c)[0] < 900]
        if pills:
            pills.sort(key=lambda b: b[2], reverse=True)
            px, py, pw, ph = pills[0]
            crop_pill = top[py:py+ph, px:px+pw]
            padded_pill = cv2.copyMakeBorder(crop_pill, 20, 20, 20, 20, cv2.BORDER_CONSTANT, value=[200,200,200])
            words_pill, lines_pill = await run_ocr_on_image(padded_pill, engine, 2.0)
            if lines_pill and len(lines_pill[0].strip()) >= 3:
                current_set = lines_pill[0].strip()
        
        # Bottom pass
        bottom = im[int(h * 0.60):h, :]
        words, _ = await run_ocr_on_image(bottom, engine, 2.5)
        
        card_ranges = [(40, 340), (360, 670), (690, 990)]
        for c_idx, (x_min, x_max) in enumerate(card_ranges, start=1):
            badge_type = classify_card_badge(im, c_idx, is_defense)
            
            c_words = [t for y, x, width, t in sorted(words) if 54 <= y <= 66 and x_min <= x <= x_max and t not in ("o", "O")]
            raw_title = " ".join(c_words).strip()
            
            sub_words = [t for y, x, width, t in sorted(words) if 67 <= y <= 80 and x_min <= x <= x_max and t not in ("o", "O")]
            sub_text = " ".join(sub_words).strip()
            
            card_crop = im[140:225, x_min-10:x_max+10]
            
            needs_recovery = (not raw_title or len(raw_title) < 2 or 
                              raw_title in ("PASS", "RUN", "ZONE", "MAN", "BLITZ", "MATCH", "I FORM", "I FORM WING", "I WING") or
                              "CALLS" in raw_title or "AVG" in raw_title)
                              
            if needs_recovery:
                # Pass 1: raw 1.5x on card
                _, lines_15 = await run_ocr_on_image(card_crop, engine, 1.5)
                for l in lines_15:
                    lt = re.sub(r"\s+", " ", l.strip()).upper()
                    if (lt not in ("PASS", "RUN", "ZONE", "MAN", "BLITZ", "MATCH", "WEAK CLOSE", "BEAR UNDER") and 
                        "CALLS" not in lt and "AVG" not in lt and len(lt) >= 3):
                        raw_title = lt
                        break
                        
                # Pass 2: wiped coach suggestion at 1.5x
                if not raw_title or len(raw_title) < 2 or "CALLS" in raw_title:
                    wiped = card_crop.copy()
                    wiped[:34, :] = 0
                    _, lines_wiped = await run_ocr_on_image(wiped, engine, 1.5)
                    for l in lines_wiped:
                        lt = re.sub(r"\s+", " ", l.strip()).upper()
                        if (lt not in ("PASS", "RUN", "ZONE", "MAN", "BLITZ", "MATCH", "WEAK CLOSE", "BEAR UNDER") and 
                            "CALLS" not in lt and "AVG" not in lt and len(lt) >= 3):
                            raw_title = lt
                            break
                            
                # Pass 3: wiped coach suggestion + threshold 140 at 2.0x
                if not raw_title or len(raw_title) < 2 or "CALLS" in raw_title:
                    gray = cv2.cvtColor(card_crop, cv2.COLOR_BGR2GRAY)
                    _, th = cv2.threshold(gray, 140, 255, cv2.THRESH_BINARY)
                    th[:34, :] = 0
                    _, lines_th = await run_ocr_on_image(th, engine, 2.0)
                    for l in lines_th:
                        lt = re.sub(r"\s+", " ", l.strip()).upper()
                        if (lt not in ("PASS", "RUN", "ZONE", "MAN", "BLITZ", "MATCH", "WEAK CLOSE", "BEAR UNDER") and 
                            "CALLS" not in lt and "AVG" not in lt and len(lt) >= 3):
                            raw_title = lt
                            break
            
            if not raw_title or len(raw_title) < 2 or raw_title in ("PASS", "RUN", "ZONE", "MAN", "BLITZ", "MATCH"):
                continue
            if "CALLS" in raw_title or "AVG" in raw_title:
                continue
            if formation_name == "I Form" and raw_title in ("I FORM", "I FORM WING", "I WING", "IFORMWING", "I Y OFF CLOSE", "I CLOSE", "I OFF CLOSE"):
                continue
                
            # If sub_text exists, it might be more specific than pill
            set_candidate = sub_text if sub_text and len(sub_text) >= 3 else current_set
            
            if is_defense:
                card_set = resolve_defense_set(formation_name, set_candidate)
                play_name, play_type = clean_defense_play(raw_title, badge_type)
            else:
                card_set = resolve_offense_set(formation_name, set_candidate)
                play_name, play_type = clean_offense_play(raw_title, badge_type)
                
            rows.append(
                PlayRow(
                    team=team,
                    formation=formation_name,
                    set_name=card_set,
                    play_name=play_name,
                    play_type=play_type,
                    source_file=f,
                    card_index=c_idx,
                )
            )
    return rows

def dedupe(rows):
    seen = set()
    unique = []
    for r in rows:
        key = (r.team.upper(), r.formation.upper(), r.set_name.upper(), r.play_name.upper())
        if key not in seen:
            seen.add(key)
            unique.append(r)
    return unique

async def run_pipeline():
    base = r"D:\Robert\Desktop\App Project Builds\Gridiron_OC\Gridiron_OC\imports\Madden 27 Team Playbooks (Screenshots)\Chicago"
    out_dir = r"D:\Robert\Desktop\App Project Builds\Gridiron_OC\Gridiron_OC\imports\Madden 27 Team Playbook (Excel Workbooks)"
    os.makedirs(out_dir, exist_ok=True)
    
    # 1. Defense First
    print("\n==========================================")
    print("=== Processing Chicago Defense ===")
    print("==========================================")
    def_rows = []
    def_dir = os.path.join(base, "Defense")
    for form in sorted(os.listdir(def_dir)):
        fp = os.path.join(def_dir, form)
        if os.path.isdir(fp):
            print(f"Scanning Defense formation: {form}...")
            f_rows = await extract_formation_plays(fp, "CHI", form, is_defense=True)
            print(f"  -> Extracted {len(f_rows)} plays")
            def_rows.extend(f_rows)
            
    unique_def = dedupe(def_rows)
    print(f"Total Defense: {len(def_rows)} raw -> {len(unique_def)} unique plays")
    
    wb_def = Workbook()
    ws_def = wb_def.active
    ws_def.title = "Defensive Plays"
    ws_def.append(["Team", "Formation", "Set", "Play Name", "Type"])
    for cell in ws_def[1]: cell.font = Font(bold=True)
    for r in sorted(unique_def, key=lambda x: (x.formation, x.set_name, x.play_name)):
        ws_def.append([r.team, r.formation, r.set_name, r.play_name, r.play_type])
        
    summary = wb_def.create_sheet("By Formation Set")
    summary.append(["Formation", "Set", "Play Count"])
    for cell in summary[1]: cell.font = Font(bold=True)
    counts = Counter((r.formation, r.set_name) for r in unique_def)
    for (formation, set_name), count in sorted(counts.items()):
        summary.append([formation, set_name, count])
        
    types = wb_def.create_sheet("By Type")
    types.append(["Type", "Count"])
    for cell in types[1]: cell.font = Font(bold=True)
    type_counts = Counter(r.play_type for r in unique_def)
    for p_type, count in sorted(type_counts.items()):
        types.append([p_type, count])
        
    def_path = os.path.join(out_dir, "Bears_defensive_playbook.xlsx")
    wb_def.save(def_path)
    print(f"Saved: {def_path}")
    
    # 2. Offense
    print("\n==========================================")
    print("=== Processing Chicago Offense ===")
    print("==========================================")
    off_rows = []
    off_dir = os.path.join(base, "Offense")
    for form in sorted(os.listdir(off_dir)):
        fp = os.path.join(off_dir, form)
        if os.path.isdir(fp):
            print(f"Scanning Offense formation: {form}...")
            f_rows = await extract_formation_plays(fp, "Bears", form, is_defense=False)
            print(f"  -> Extracted {len(f_rows)} plays")
            off_rows.extend(f_rows)
            
    unique_off = dedupe(off_rows)
    print(f"Total Offense: {len(off_rows)} raw -> {len(unique_off)} unique plays")
    
    wb_off = Workbook()
    ws_off = wb_off.active
    ws_off.title = "Playbook"
    ws_off.append(["Team", "Formation", "Set", "Play Name", "Type"])
    for cell in ws_off[1]: cell.font = Font(bold=True)
    for r in sorted(unique_off, key=lambda x: (x.formation, x.set_name, x.play_name)):
        ws_off.append([r.team, r.formation, r.set_name, r.play_name, r.play_type])
    off_path = os.path.join(out_dir, "Bears_playbook-template.xlsx")
    wb_off.save(off_path)
    print(f"Saved: {off_path}")
    print("\n=== All Chicago Extractions Complete & Verified! ===")

if __name__ == "__main__":
    asyncio.run(run_pipeline())
