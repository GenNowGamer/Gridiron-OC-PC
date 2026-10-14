"""Write Madden 27 play JSON in the same shape Gridiron OC uses for Madden 26."""
import json
from openpyxl import load_workbook

WORKBOOK = r"D:\Robert\Desktop\App Project Builds\Gridiron_OC\Gridiron_OC\output_madden27_playbooks.xlsx"
OUT_OFF = r"D:\Robert\Desktop\App Project Builds\Gridiron_OC\Gridiron_OC\PC\data\madden27\plays.json"
OUT_DEF = r"D:\Robert\Desktop\App Project Builds\Gridiron_OC\Gridiron_OC\PC\data\madden27\defensive_plays.json"

TEAM_CODES = {
    "49ers": "SF", "Bears": "CHI", "Bengals": "CIN", "Bills": "BUF", "Broncos": "DEN",
    "Browns": "CLE", "Buccaneers": "TB", "Cardinals": "ARI", "Chargers": "LAC",
    "Chiefs": "KC", "Colts": "IND", "Commanders": "WAS", "Cowboys": "DAL",
    "Dolphins": "MIA", "Eagles": "PHI", "Falcons": "ATL", "Giants": "NYG",
    "Jaguars": "JAX", "Jets": "NYJ", "Lions": "DET", "Packers": "GB", "Panthers": "CAR",
    "Patriots": "NE", "Raiders": "LV", "Rams": "LAR", "Ravens": "BAL", "Saints": "NO",
    "Seahawks": "SEA", "Steelers": "PIT", "Texans": "HOU", "Titans": "TEN", "Vikings": "MIN",
}
OFFENSE_TYPES = {"Run": "RUN", "Pass": "PASS", "Play Action": "PA", "RPO": "RPO"}
DEFENSE_TYPES = {"MAN": "MAN", "ZONE": "ZONE", "BLITZ": "BLITZ", "MATCH": "MATCH"}
DEFENSE_CONCEPT = {"MAN": "Man", "ZONE": "Zone", "BLITZ": "Blitz", "MATCH": "Match"}
DEFENSE_PHRASES = (
    ("tampa 2", "Tampa 2"), ("cover 0", "Cover 0"), ("cover 1", "Cover 1"),
    ("cover 2", "Cover 2"), ("cover 3", "Cover 3"), ("cover 4", "Cover 4"),
    ("cover 6", "Cover 6"), ("cover 9", "Cover 9"), ("goal line", "Goal Line"),
    ("press", "Press"), ("sky", "Sky"), ("buzz", "Buzz"), ("fire", "Fire"),
    ("flat", "Flat"), ("robber", "Robber"), ("bracket", "Bracket"),
    ("overload", "Overload"), ("pinch", "Pinch"), ("cloud", "Cloud"),
    ("seam", "Seam"), ("loop", "Loop"), ("contain", "Contain"), ("storm", "Storm"),
    ("spy", "Spy"), ("mug", "Mug"), ("twist", "Twist"),
)


def text(value):
    return " ".join(str(value or "").split())


def offense_concepts(play_type, play_name, formation, set_name):
    blob = f"{play_name} {set_name} {formation}".lower()
    tags = []
    if play_type == "RPO":
        tags.append("RPO")
    if play_type == "PA":
        tags.append("Play Action")
    checks = (
        ("inside zone", "Inside Zone"), ("split zone", "Split Zone"),
        (("outside zone", "wide zone"), "Wide Zone"), ("duo", "Duo"),
        ("counter", "Counter"), ("trap", "Trap"), ("power", "Power O"),
        ("draw", "Draw"), (("dive", "plunge"), "Dive"), ("screen", "Screen"),
        ("stick", "Stick"), ("spacing", "Spacing"), (("slant", "slants"), "Slant"),
        ("mesh", "Mesh"), ("drag", "Drag"), ("under", "Under"), ("cross", "Cross"),
        ("over", "Deep Over"), ("dagger", "Dagger"), ("post", "Post"),
        ("corner", "Corner"), ("bench", "Bench"), (("flood", "sail"), "Flood"),
        (("boot", "sprint"), "Boot"), ("leak", "Leak"), ("wheel", "Wheel"),
        ("seam", "Seam"), ("curl", "Curl"), ("flat", "Flat"), ("swing", "Swing"),
        (("out", "outs"), "Out"),
    )
    for needle, label in checks:
        needles = needle if isinstance(needle, tuple) else (needle,)
        if any(item in blob for item in needles):
            tags.append(label)
    if any(item in blob for item in ("stick", "spacing", "bench", "slant", "under", "curl", "outs")):
        tags.append("Quick Game")
    if play_type == "RUN" and not tags:
        tags.append("Run")
    if play_type == "PASS" and not tags:
        tags.append("Pass")
    return list(dict.fromkeys(tags))


def offense_strong_vs(play_type, play_name, concepts):
    text_value = play_name.lower()
    concept_text = [item.lower() for item in concepts]

    def has(*terms):
        return any(term in text_value or any(term in concept for concept in concept_text) for term in terms)

    strong = []
    if has("mesh", "drag", "slant", "cross", "wheel", "leak", "under", "dig"):
        strong.extend(["Man", "Cover 1", "Cover 6"])
    if has("seam", "post", "dagger"):
        strong.extend(["Cover 1", "Tampa 2"])
    if has("flood", "corner", "bench", "spacing", "stick", "quick game", "out"):
        strong.extend(["Cover 2", "Cover 3", "Tampa 2"])
    if has("dagger", "post", "seam", "deep over", "scissors", "shot"):
        strong.extend(["Cover 4", "Cover 6"])
    if has("boot", "play action", "screen", "rpo", "flat", "swing"):
        strong.append("Cover 9")
    if play_type == "RUN" or has("screen", "quick game", "trap", "inside zone", "split zone", "dive", "leak"):
        strong.extend(["Blitz", "Cover 9"])
    return list(dict.fromkeys(strong))


def defense_concepts(play_type, play_name, formation, set_name):
    blob = f"{play_name} {formation} {set_name}".lower()
    tags = [DEFENSE_CONCEPT[play_type]]
    for needle, label in DEFENSE_PHRASES:
        if needle in blob and label not in tags:
            tags.append(label)
    if "invert" in blob and "Cover 2" not in tags:
        tags.append("Cover 2")
    return tags


def play_record(team, formation, set_name, play_name, play_type, concepts, strong_vs):
    return {
        "team": team,
        "formation": formation,
        "set": set_name,
        "play_name": play_name,
        "type": play_type,
        "concepts": concepts,
        "strong_vs": strong_vs,
        "situations": [],
    }


def main():
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    offense = []
    defense = []
    for row in workbook["Offense"].iter_rows(min_row=2, values_only=True):
        team = TEAM_CODES[row[0]]
        formation, set_name, play_name = text(row[1]), text(row[2]), text(row[3])
        play_type = OFFENSE_TYPES[text(row[4])]
        concepts = offense_concepts(play_type, play_name, formation, set_name)
        offense.append(play_record(
            team, formation, set_name, play_name, play_type, concepts,
            offense_strong_vs(play_type, play_name, concepts),
        ))
    for row in workbook["Defense"].iter_rows(min_row=2, values_only=True):
        team = TEAM_CODES[row[0]]
        formation, set_name, play_name = text(row[1]), text(row[2]), text(row[3])
        play_type = DEFENSE_TYPES[text(row[4]).upper()]
        defense.append(play_record(
            team, formation, set_name, play_name, play_type,
            defense_concepts(play_type, play_name, formation, set_name), [],
        ))
    for path, rows in ((OUT_OFF, offense), (OUT_DEF, defense)):
        with open(path, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(rows, handle, indent=2, ensure_ascii=False)
            handle.write("\n")
    print("offense", len(offense), "teams", len({row["team"] for row in offense}))
    print("defense", len(defense), "teams", len({row["team"] for row in defense}))


if __name__ == "__main__":
    main()
