import os
import sys
import re
import glob
import asyncio
import difflib
import unicodedata
import cv2
import numpy as np
from collections import Counter
from dataclasses import dataclass
from openpyxl import Workbook
from openpyxl.styles import Font

import winsdk.windows.graphics.imaging as imaging
import winsdk.windows.media.ocr as ocr
import winsdk.windows.storage as storage

import catalog_plays
import catalog_sets

@dataclass
class PlayRow:
    team: str
    formation: str
    set_name: str
    play_name: str
    play_type: str
    source_file: str
    card_index: int

DEF_REP_EXTRA = {
    # 4-3 & general defense corrections
    "SAMISTING": ("Sam 1 Sting", "BLITZ"),
    "SSBUTZI": ("SS Blitz 1", "BLITZ"),
    "SAWOBUST": ("Sam 0 Blast", "BLITZ"),
    "HAMMEROBLAST": ("Hammer 0 Blast", "BLITZ"),
    "COVER8SKYWK": ("Cover 3 Sky Wk", "ZONE"),
    "COVER21WERT": ("Cover 2 Invert", "ZONE"),
    "21WERTHARDFLAT": ("2 Invert Hard Flat", "ZONE"),
    "21NVERTHARDFLAT": ("2 Invert Hard Flat", "ZONE"),
    "COVER8HARDFLAT": ("Cover 3 Hard Flat", "ZONE"),
    "COVER81URDFLAT": ("Cover 3 Hard Flat", "ZONE"),
    "COVER8MATCH": ("Cover 3 Match", "MATCH"),
    "COVER8SKY": ("Cover 3 Sky", "ZONE"),
    "COVER4WARTERS": ("Cover 4 Quarters", "MATCH"),
    "COVER4PRESS": ("Cover 4 Press", "MATCH"),
    "COVERIHOLE": ("Cover 1 Hole", "MAN"),
    "IROBBERPRESS": ("1 Robber Press", "MAN"),
    "IHOLEPRESS": ("1 Hole Press", "MAN"),
    "ICOUTAINPRESS": ("1 Contain Press", "MAN"),
    "COVERICONUIN": ("Cover 1 Contain", "MAN"),
    "COVER2MAN": ("Cover 2 Man", "MAN"),
    "COVER2HARDFLAT": ("Cover 2 Hard Flat", "ZONE"),
    "LBSTINGI": ("LB Sting 1", "BLITZ"),
    "OLBFIREMAN": ("OLB Fire Man", "BLITZ"),
    "SAMEDGE8": ("Sam Edge 3", "BLITZ"),
    "THUNDERSMOKE": ("Thunder Smoke", "BLITZ"),
    "DBLSAFETYBLITZ": ("Dbl Safety Blitz", "BLITZ"),
    "FREEFIRE": ("Free Fire", "BLITZ"),
    "FREEFIREO": ("Free Fire 0", "BLITZ"),
    "FSFLRE8": ("FS Fire 3", "BLITZ"),
    "FSFIREI": ("FS Fire 1", "BLITZ"),
    "SAMCUSH8": ("Sam Crash 3", "BLITZ"),
    "SAMCRASHI": ("Sam Crash 1", "BLITZ"),
    "WILLGOFLRE8": ("Will Go Fire 3", "BLITZ"),
    "PINCHBIMO": ("Pinch Buck 0", "BLITZ"),
    "PINCHDOG2PRESS": ("Pinch Dog 2 Press", "BLITZ"),
    "STORMBUVEI": ("Storm Brave 1", "BLITZ"),
    "WILLBUCK3PRESS": ("Will Buck 3 Press", "BLITZ"),
    "CROSS3FIRE": ("Cross 3 Fire", "BLITZ"),
    "SS2TRAP": ("SS 2 Trap", "BLITZ"),
    "SAWBUTZI": ("Saw Blitz 3", "BLITZ"),
    "COVERS": ("Cover 6", "MATCH"),
    "COVERG": ("Cover 6", "MATCH"),
    "TAMPA2": ("Tampa 2", "ZONE"),
    "TAUPA2": ("Tampa 2", "ZONE"),
    "COVER8CWJDSTR": ("Cover 3 Cloud Str", "ZONE"),

    # 46
    "COVER8": ("Cover 3", "ZONE"),
    "COVERI": ("Cover 1", "BLITZ"),
    "EUGAGEEIGHT": ("Engage Eight", "BLITZ"),
    "ENGAGEEIGHT": ("Engage Eight", "BLITZ"),
    "GAPPRSS": ("Gap Press", "BLITZ"),
    "LBDOGS8": ("LB Dogs 3", "BLITZ"),
    "LBDOGS": ("LB Dogs", "BLITZ"),

    # Dollar
    "ZEROBUTZ": ("Zero Blitz", "BLITZ"),
    "COVER2DROP": ("Cover 2 Drop", "ZONE"),
    "DBFIRE2": ("DB Fire 2", "BLITZ"),
    "DBFLRE2": ("DB Fire 2", "BLITZ"),
    "OVERLOAD8SEAM": ("Overload 3 Seam", "BLITZ"),
    "OVERLOAD8SUM": ("Overload 3 Seam", "BLITZ"),
    "OVERLOADBLITZ": ("Overload Blitz", "BLITZ"),
    "COVER4DROP": ("Cover 4 Drop", "ZONE"),
    "DTMIKELOOP0": ("DT Mike Loop 0", "BLITZ"),
    "DTMIKELOOPO": ("DT Mike Loop 0", "BLITZ"),
    "DTMIKELOOP1": ("DT Mike Loop 1", "BLITZ"),
    "DTMIKELOOPI": ("DT Mike Loop 1", "BLITZ"),
    "ERIROBBERPRESS": ("Cover 1 Robber Press", "MAN"),
    "COVERIOBERPRESS": ("Cover 1 Robber Press", "MAN"),

    # Goal Line
    "GOPINCH": ("60 Pinch", "BLITZ"),
    "60PINCH": ("60 Pinch", "BLITZ"),
    "GOBASE": ("60 Base", "BLITZ"),
    "60BASE": ("60 Base", "BLITZ"),
    "GOHALFOUT": ("60 Half Out", "BLITZ"),
    "60HALFOUT": ("60 Half Out", "BLITZ"),
    "GOHA1FUJT": ("60 Half Out", "BLITZ"),
    "GOH1FUJT": ("60 Half Out", "BLITZ"),
    "GOH1FLT": ("60 Half Out", "BLITZ"),
    "GOH1FLUT": ("60 Half Out", "BLITZ"),
    "GOH1FUT": ("60 Half Out", "BLITZ"),
    "GOHAIFWT": ("60 Half Out", "BLITZ"),
    "GOHALFWT": ("60 Half Out", "BLITZ"),
    "GOOUT": ("60 Out", "BLITZ"),
    "60OUT": ("60 Out", "BLITZ"),
    "EOMN": ("60 Out", "BLITZ"),
    "EOMUT": ("60 Out", "BLITZ"),
    "EOMT": ("60 Out", "BLITZ"),
    "GOOUTJACKS": ("60 Out Jacks", "BLITZ"),
    "60OUTJACKS": ("60 Out Jacks", "BLITZ"),
    "GOMJTJACKS": ("60 Out Jacks", "BLITZ"),
    "GUTS": ("Guts", "BLITZ"),
    "UTS": ("Guts", "BLITZ"),
    "UIS": ("Guts", "BLITZ"),
    "NJTS": ("Guts", "BLITZ"),
    "MNS": ("Guts", "BLITZ"),
    "WTS": ("Guts", "BLITZ"),
    "GWTS": ("Guts", "BLITZ"),
    "OHUFDUT": ("60 Half Out", "BLITZ"),
    "GLMAN": ("GL Man", "BLITZ"),
    "MLBGAPA": ("MLB Gap A", "BLITZ"),
    "GAPSAU": ("Gaps All", "BLITZ"),
    "GAPSALL": ("Gaps All", "BLITZ"),
    "GAPSAB": ("Gaps A B", "BLITZ"),
    "GAPSB": ("Gaps B", "BLITZ"),
    "GAPSAULB": ("Gaps All LB", "BLITZ"),
    "GAPSALLLB": ("Gaps All LB", "BLITZ"),
    "GOALLINEPINCH": ("60 Pinch", "BLITZ"),

    # Prevent
    "MAN8DEEP": ("Man 3 Deep", "MAN"),
    "MAN3DEEP": ("Man 3 Deep", "MAN"),
    "PREVENT": ("Prevent", "ZONE"),
    "COVER2DO": ("Cover 2 Drop", "ZONE"),
    "8DEEPBUTZ": ("3 Deep Blitz", "BLITZ"),
    "8DEEPBLITZ": ("3 Deep Blitz", "BLITZ"),
    "08DEEPBUTZ": ("3 Deep Blitz", "BLITZ"),
    "08DEEPBLITZ": ("3 Deep Blitz", "BLITZ"),
    "IDOUBLEBUZZ": ("3 Double Buzz", "ZONE"),

    # 3-3-5
    "PINCHBUTZO": ("Pinch Blitz 0", "BLITZ"),
    "PINCHBLITZO": ("Pinch Blitz 0", "BLITZ"),
    "COVER4URTERS": ("Cover 4 Quarters", "MATCH"),
    "COVEUURTERS": ("Cover 4 Quarters", "MATCH"),
    "SAWBUTZ3": ("Saw Blitz 3", "BLITZ"),
    "SAWBUTZ8": ("Saw Blitz 3", "BLITZ"),
    "COVERHOLEI": ("Cover Hole 1", "MAN"),
    "COVER4PILUS": ("Cover 4 Palms", "MATCH"),
    "COVER4US": ("Cover 4 Palms", "MATCH"),
    "HOTBLLTZ8": ("Hot Blitz 3", "BLITZ"),
    "COVERIOBER": ("Cover 1 Robber", "MAN"),
    "CORNERBUTZ3": ("Corner Blitz 3", "BLITZ"),
    "DMJBLEBUCKET": ("Double Bracket", "MAN"),
    "DUIEBUCKET": ("Double Bracket", "MAN"),
    "SLOTBUTZ3": ("Slot Blitz 3", "BLITZ"),
    "SLOTBUTZ8": ("Slot Blitz 3", "BLITZ"),
    "COVERGWILLIE": ("Cover 6 Willie", "MATCH"),
    "COVER8BLRZ": ("Cover 3 Buzz", "ZONE"),
    "COVER8BTM": ("Cover 3 Buzz", "ZONE"),
    "SAMMIKE8PRESS": ("Sam Mike 3 Press", "BLITZ"),
    "MIKEBLITZI": ("Mike Blitz 1", "BLITZ"),
    "MIKEBLITZJ": ("Mike Blitz 1", "BLITZ"),
    "CORUERBLITZ2": ("Corner Blitz 2", "BLITZ"),
    "CORNERBLITZ2": ("Corner Blitz 2", "BLITZ"),

    # Dime
    "EDGEBUTZO": ("Edge Blitz 0", "BLITZ"),
    "EDGEBLITZ0": ("Edge Blitz 0", "BLITZ"),
    "COVER8BLUMATCH": ("Cover 3 Buzz Match", "MATCH"),
    "COVER8BLRZMATCH": ("Cover 3 Buzz Match", "MATCH"),
    "IDOUBLEWRI": ("1 Double WR1", "MAN"),
    "FIELDSIM8": ("Field Sim 3", "BLITZ"),
    "FIELDSM8": ("Field Sim 3", "BLITZ"),
    "DMJBLEBUCKETSWITCH": ("Double Bracket Switch", "MAN"),
    "COVERIPALMS": ("Cover 4 Palms", "MATCH"),
    "COVERDPALMS": ("Cover 4 Palms", "MATCH"),
    "OVERSTORMBRAVE": ("Over Storm Brave", "BLITZ"),
    "COVERIROBBER": ("Cover 1 Robber", "MAN"),
    "EDGEBLITZ8": ("Edge Blitz 3", "BLITZ"),
    "EDGEBLITZ3": ("Edge Blitz 3", "BLITZ"),
    "COVER6SHW2": ("Cover 6 Show 2", "MATCH"),
    "COVER6SHW12": ("Cover 6 Show 2", "MATCH"),
    "OVERLOAD8PRESS": ("Overload 3 Press", "BLITZ"),
    "COVER8SHOW2": ("Cover 9 Show 2", "MATCH"),
    "COVER9SHOW2": ("Cover 9 Show 2", "MATCH"),
    "NICKEL2TRAP": ("Nickel 2 Trap", "BLITZ"),
    "NICKEL2TUP": ("Nickel 2 Trap", "BLITZ"),
    "NICKEL2TRP": ("Nickel 2 Trap", "BLITZ"),
    "NICKELSLM2": ("Nickel Sim 2", "BLITZ"),
    "NICKELS1M2": ("Nickel Sim 2", "BLITZ"),
    "COVERIHDLE": ("Cover 1 Hole", "MAN"),
    "SIMPRESSURE8": ("Sim Pressure 3", "BLITZ"),
    "FSWILLBLITZO": ("FS Will Blitz 0", "BLITZ"),
    "BUCKET9SWLTCH": ("Bracket 9 Switch", "MAN"),
    "BRACKET9SWLTCH": ("Bracket 9 Switch", "MAN"),
    "COVER2MATCH": ("Cover 2 Match", "MATCH"),
    "COVER8SAMMLE": ("Cover 3 Seam", "MATCH"),
    "COVERSAMMIE": ("Cover 3 Seam", "MATCH"),
    "BUCKETSWITCHWILLIE": ("Bracket Switch Willie", "MAN"),
    "BUCKETSWITCHWILUE": ("Bracket Switch Willie", "MAN"),
    "IDOUBLESLOT": ("1 Double Slot", "MAN"),
    "IDOUBLESLDT": ("1 Double Slot", "MAN"),
    "1DOUBLETE": ("1 Double TE", "MAN"),
    "IDOUBLETE": ("1 Double TE", "MAN"),
    "IDOUBLEWR2": ("1 Double WR2", "MAN"),
    "JDOUBLEWR2": ("1 Double WR2", "MAN"),
    "ROUCOVER21NVERT": ("Roll Cover 2 Invert", "ZONE"),
    "ROLLCOVER21NVERT": ("Roll Cover 2 Invert", "ZONE"),
    "ROLL28CLWD": ("Roll 2 Cloud", "ZONE"),
    "MLL28CWJD": ("Roll 2 Cloud", "ZONE"),
    "NL28CLOUD": ("Roll 2 Cloud", "ZONE"),
    "NL28CLOLJD": ("Roll 2 Cloud", "ZONE"),
    "MLLCOVER8CLOLD": ("Roll Cover 3 Cloud", "ZONE"),
    "ROLLCOVER8C": ("Roll Cover 3 Cloud", "ZONE"),
    "ROLLCOVER8CLOLD": ("Roll Cover 3 Cloud", "ZONE"),
    "COVER8CLOUD": ("Cover 3 Cloud", "ZONE"),
    "COVER8CIDJD": ("Cover 3 Cloud", "ZONE"),
    "COVER8CIMJD": ("Cover 3 Cloud", "ZONE"),
    "COVER8CLMJD": ("Cover 3 Cloud", "ZONE"),
    "MISIMPRESSURE": ("Mug Sim Pressure", "BLITZ"),
    "MUGDTLOOPI": ("Mug DT Loop 1", "BLITZ"),
    "MLSSBLLTZTEX8": ("Mug Blitz Tex 3", "BLITZ"),
    "MTMLLTZTEX8": ("Mug Blitz Tex 3", "BLITZ"),
    "BLITZDB": ("DB Blitz", "BLITZ"),
    "DIMEBLITZ2": ("Dime Blitz 2", "BLITZ"),
    "FSMLDDLEBLLTZ8": ("FS Middle Blitz 3", "BLITZ"),
    "COVER8BUMATCH": ("Cover 3 Buzz Match", "MATCH"),
    "COVER2ROU": ("Cover 2 Roll", "ZONE"),
    "SSSIMPRESSURE8": ("SS Sim Pressure 3", "BLITZ"),
    "FOXFIREZONEPRESS": ("Fox Fire Zone Press", "BLITZ"),
    "COVERICONTAINSPY": ("Cover 1 Contain Spy", "MAN"),
    "SSOVERLOAD8": ("SS Overload 3", "BLITZ"),

    # Nickel
    "CWER4PALUS": ("Cover 4 Palms", "MATCH"),
    "COVER4MLUS": ("Cover 4 Palms", "MATCH"),
    "FSWIUBUTZO": ("FS Will Blitz 0", "BLITZ"),
    "8BUZZMATCHWK": ("Cov 3 Buzz Match Wk", "MATCH"),
    "COV8BUZZMATCHWK": ("Cov 3 Buzz Match Wk", "MATCH"),
    "MBBERPRSS": ("1 Robber Press", "MAN"),
    "NICKELBUTZ8": ("Nickel Blitz 3", "BLITZ"),
    "NICKELBLLTZ8": ("Nickel Blitz 3", "BLITZ"),
    "LBBUTZI": ("LB Blitz 1", "BLITZ"),
    "SHWJSSSLM2": ("Show SS Sim 2", "BLITZ"),
    "SHMNSSSLM2": ("Show SS Sim 2", "BLITZ"),
    "2INVERTLURK": ("Cover 2 Invert Lurk", "ZONE"),
    "21MVERTLURK": ("Cover 2 Invert Lurk", "ZONE"),
    "MIDBLITZO": ("Mid Blitz 0", "BLITZ"),
    "BLITZLOOP3": ("Blitz Loop 3", "BLITZ"),
    "BLITZLOOP8": ("Blitz Loop 3", "BLITZ"),
    "COVER8SEAU": ("Cover 3 Seam", "MATCH"),
    "ER8SEAU": ("Cover 3 Seam", "MATCH"),
    "NICKELOVERLOAD3": ("Nickel Overload 3", "BLITZ"),
    "NICKELDVERLOD3": ("Nickel Overload 3", "BLITZ"),
    "LBBUTZ8": ("LB Blitz 3", "BLITZ"),
    "SSBLLTZ8": ("SS Blitz 3", "BLITZ"),
    "NICKELDOG8BUZZ": ("Nickel Dog 3 Buzz", "BLITZ"),
    "NICKELBLITZO": ("Nickel Blitz 0", "BLITZ"),
    "LBPOPO": ("LB Pop 0", "BLITZ"),
    "LBPOP0": ("LB Pop 0", "BLITZ"),
    "BENCHBLITZ3": ("Bench Blitz 3", "BLITZ"),
    "BENCHBLLTZ8": ("Bench Blitz 3", "BLITZ"),
    "LBDETWIST3": ("LB DE Twist 3", "BLITZ"),
    "LBDETWLST8": ("LB DE Twist 3", "BLITZ"),
    "SAMMIKELOOP3": ("Sam Mike Loop 3", "BLITZ"),
    "SAMMIKELOOP8": ("Sam Mike Loop 3", "BLITZ"),
    "SAMMIKELOOPI": ("Sam Mike Loop 3", "BLITZ"),
    "LBCROSS3SHW12": ("LB Cross 3 Show 2", "BLITZ"),
    "LBCROSS8SHW12": ("LB Cross 3 Show 2", "BLITZ"),
    "DTMIKELOOP3": ("DT Mike Loop 3", "BLITZ"),
    "DTMIKELODPI": ("DT Mike Loop 3", "BLITZ"),
    "DTMIKE1POP0": ("DT Mike 1 Pop 0", "BLITZ"),
    "DTMIKEPOPO": ("DT Mike 1 Pop 0", "BLITZ"),
    "HOTBLLTZ3": ("Hot Blitz 3", "BLITZ"),
    "IDTLBLODP": ("1 DT LB Loop", "BLITZ"),
    "COVERIGAMS": ("Cover 1 Games", "BLITZ"),
    "BLITZLOOPO": ("Blitz Loop 0", "BLITZ"),
    "BLITZLOOPSM8": ("Blitz Loop Sim 3", "BLITZ"),
    "REDZONEDTDO": ("Redzone DT Drop", "ZONE"),
    "SSDOG8": ("SS Dog 3", "BLITZ"),
    "SSSIM2": ("Cross Sim 2", "BLITZ"),
    "CBBENCHSLM2": ("CB Bench Sim 2", "BLITZ"),
    "ROLL28CLOUD": ("Roll 2 Cloud", "ZONE"),
    "ROLLCOVER8CLOUD": ("Roll Cover 3 Cloud", "ZONE"),
    "2-4LODDALMUG": ("2-4 Load Dbl Mug", "ZONE"),

    # Atlanta & general new corrections
    "COMAWARTERS": ("Cover 4 Quarters", "MATCH"),
    "CWERIURTERS": ("Cover 4 Quarters", "MATCH"),
    "CWER4URTERS": ("Cover 4 Quarters", "MATCH"),
    "CWER4QTURTERS": ("Cover 4 Quarters", "MATCH"),
    "COVERIWARTERS": ("Cover 4 Quarters", "MATCH"),
    "COME": ("Cover 6", "MATCH"),
    "COMS": ("Cover 6", "MATCH"),
    "COVNG": ("Cover 6", "MATCH"),
    "COVNS": ("Cover 6", "MATCH"),
    "IUVERT2SLM": ("Invert 2 Sim", "BLITZ"),
    "INVERT2SLM": ("Invert 2 Sim", "BLITZ"),
    "INVERT2SIM": ("Invert 2 Sim", "BLITZ"),
    "WILLBUTZ8": ("Will Blitz 3", "BLITZ"),
    "WIUBUTZ8": ("Will Blitz 3", "BLITZ"),
    "WULBUTZ8": ("Will Blitz 3", "BLITZ"),
    "UIKEBUTZI": ("Mike Blitz 1", "BLITZ"),
    "UIIEBUTZI": ("Mike Blitz 1", "BLITZ"),
    "STORUBUVEI": ("Storm Brave 1", "BLITZ"),
    "CMQ8SKY": ("Cover 3 Sky", "ZONE"),
    "C8SKY": ("Cover 3 Sky", "ZONE"),
    "CSLM2": ("Nickel Sim 2", "BLITZ"),
    "CSIM2": ("Nickel Sim 2", "BLITZ"),
    "C8BLZ": ("Cover 3 Buzz", "ZONE"),
    "CIPALMS": ("Cover 4 Palms", "MATCH"),
    "COVNAPALMS": ("Cover 4 Palms", "MATCH"),
    "COVN3SKY": ("Cover 3 Sky", "ZONE"),
    "COVN3MATCH": ("Cover 3 Match", "MATCH"),
    "CN3BUZZ": ("Cover 3 Buzz", "ZONE"),
    "SANUIKE8PRES": ("Sam Mike 3 Press", "BLITZ"),
    "CORNERBUTZ8": ("Corner Blitz 3", "BLITZ"),
    "DOIIEBRACKET": ("Double Bracket", "MAN"),
    "CN21NVERT": ("Cover 2 Invert", "ZONE"),
    "CM21NVERT": ("Cover 2 Invert", "ZONE"),
    "CNGSHOW2": ("Cover 9 Show 2", "MATCH"),
    "COVER8SHN2": ("Cover 9 Show 2", "MATCH"),
    "IDWLETE": ("1 Double TE", "MAN"),
    "1DOJBLETE": ("1 Double TE", "MAN"),
    "IDLEWRI": ("1 Double WR1", "MAN"),
    "IDULEWRI": ("1 Double WR1", "MAN"),
    "IDMJBLEWRL": ("1 Double WR1", "MAN"),
    "MLLCOVER8CLMJD": ("Roll Cover 3 Cloud", "ZONE"),
    "NLCOVER3CLMJD": ("Roll Cover 3 Cloud", "ZONE"),
    "ROILCOVER8CLMLD": ("Roll Cover 3 Cloud", "ZONE"),
    "MTL28C1DTN": ("Roll 2 Cloud", "ZONE"),
    "ROIL28CL": ("Roll 2 Cloud", "ZONE"),
    "DU28CLMJD": ("Roll 2 Cloud", "ZONE"),
    "GWJTJACKS": ("60 Out Jacks", "BLITZ"),
    "21NVERTUJRK": ("Cover 2 Invert Lurk", "ZONE"),
    "BUTZTEX8SLM8": ("Blitz Tex 3 Sim 3", "BLITZ"),
    "BUTZTEX8SIM8": ("Blitz Tex 3 Sim 3", "BLITZ"),
    "COVER9SAUMLE": ("Cover 3 Seam", "MATCH"),
    "EDCEBUTZ8": ("Edge Blitz 3", "BLITZ"),
    "IBBUTZ8": ("LB Blitz 3", "BLITZ"),
    "8DBIMABLECLOUD": ("3 Double Cloud", "ZONE"),
    "8DBIUIBLECTOUD": ("3 Double Cloud", "ZONE"),
    "8DMJBLEMABIE": ("3 Buzz Mable", "MATCH"),
    "C8MABLE": ("3 Buzz Mable", "MATCH"),
    "COVER8UBLE": ("3 Buzz Mable", "MATCH"),
    "C8SPY": ("Cover 3 Spy", "ZONE"),
    "COVER8SPY": ("Cover 3 Spy", "ZONE"),
    "COVERIDO": ("Cover 4 Drop", "ZONE"),
    "COVERIDROP": ("Cover 4 Drop", "ZONE"),
    "INDERSMOKE": ("Under Smoke", "BLITZ"),
    "INDERSUOXE": ("Under Smoke", "BLITZ"),
    "8DOUBLEWZZ": ("3 Double Buzz", "ZONE"),
    "UANUP8DEEP": ("Man Up 3 Deep", "MAN"),
    "FZ8SKY": ("FZ 3 Sky", "BLITZ"),
    "COVER2MANSPY": ("Cover 2 Man Spy", "MAN"),
    "COVERITHIEF": ("Cover 1 Thief", "MAN"),
    "OVERLOAD8": ("Overload 3", "BLITZ"),
    "OVERLOADA": ("Overload 3", "BLITZ"),
    "MIKESAMC8": ("Mike Sam Crash Press", "BLITZ"),
    "DOUBLECBBUTZ": ("Double CB Blitz", "BLITZ"),
    "DOUBLECBBLITZ": ("Double CB Blitz", "BLITZ"),
    "UNPA2": ("Tampa 2", "ZONE"),
    "CAWARTERS": ("Cover 4 Quarters", "MATCH"),
    "CAURTERS": ("Cover 4 Quarters", "MATCH"),
    "CN3CLOUDSHOW2": ("Cover 3 Cloud Show 2", "ZONE"),
    "CN3CIOUDSHOW2": ("Cover 3 Cloud Show 2", "ZONE"),
    "CN8CLOUDSHOW2": ("Cover 3 Cloud Show 2", "ZONE"),
    "CN8CIOUDSHOW2": ("Cover 3 Cloud Show 2", "ZONE"),
    "COVER8CIOUDSHM": ("Cover 3 Cloud Str", "ZONE"),
    "GOWTJACKS": ("60 Out Jacks", "BLITZ"),
    "CA8BLUZ": ("Cover 3 Buzz", "ZONE"),
    "CAIPALMS": ("Cover 4 Palms", "MATCH"),
    "CA8SKY": ("Cover 3 Sky", "ZONE"),
    "CASLM2": ("Nickel Sim 2", "BLITZ"),
    "ROIL28CLU": ("Roll 2 Cloud", "ZONE"),
    "CA8SPY": ("Cover 3 Spy", "ZONE"),

    # Dallas additions
    "CAIURTERS": ("Cover 4 Quarters", "MATCH"),
    "DULECBBUTZ": ("Double CB Blitz", "BLITZ"),
    "NIKESAMCA8": ("Mike Sam Crash Press", "BLITZ"),
    "CA8BUZZ": ("Cover 3 Buzz", "ZONE"),
    "CA8MATCH": ("Cover 3 Match", "MATCH"),
    "COVNGSXY": ("Cover 3 Sky", "ZONE"),
    "COVNGSKY": ("Cover 3 Sky", "ZONE"),
    "COMLSTLNG": ("Cover 1 Sting", "MAN"),
    "COVNAD0": ("Cover 4 Drop", "ZONE"),
    "COVNAD8": ("Cover 4 Drop", "ZONE"),
    "BUCKETBSWITCH": ("Bracket 9 Switch", "MAN"),
    "COVN21MVERT": ("Cover 2 Invert", "ZONE"),
    "COVNGSHM": ("Cover 9 Show 2", "MATCH"),
    "POUNDER": ("Pounder", "MATCH"),
    "RING": ("Ring", "MATCH"),
    "ROUKING": ("Roll King", "MATCH"),
    "DUKING": ("Roll King", "MATCH"),
    "ROILRING": ("Roll King", "MATCH"),
    "ROURING": ("Roll King", "MATCH"),
    "ROLLKING": ("Roll King", "MATCH"),
    "CA2MAN": ("Cover 2 Man", "MAN"),
    "GOJJT": ("60 Out", "BLITZ"),
    "BUCKET8SWLTCH": ("Bracket 9 Switch", "MAN"),
    "BUTZTEX8SLU8": ("Blitz Tex 3 Sim 3", "BLITZ"),
    "COFFEEHMJSE3SEAM": ("Coffeehouse 3 Seam", "BLITZ"),
    "COFFEEHOUSE8SEAM": ("Coffeehouse 3 Seam", "BLITZ"),
    "COFFEEHOUSE8SUU": ("Coffeehouse 3 Seam", "BLITZ"),
    "COFFEEHWSE3SUM": ("Coffeehouse 3 Seam", "BLITZ"),
    "COFFEEHWSE8SUM": ("Coffeehouse 3 Seam", "BLITZ"),
    "COFFEEHOUSE3SEAM": ("Coffeehouse 3 Seam", "BLITZ"),
    "COVER8ZDNE": ("Cover 3 Match", "MATCH"),
    "COVN3HUTCH": ("Cover 3 Match", "MATCH"),
    "COVNAWARTERS": ("Cover 4 Quarters", "MATCH"),
    "IBBUTZO": ("LB Blitz 0", "BLITZ"),
    "IDCIBLEWRI": ("1 Double WR1", "MAN"),
    "IDCIBLEWR1": ("1 Double WR1", "MAN"),
    "ROU28CLMN": ("Roll 2 Cloud", "ZONE"),
    "ROU28CLMN": ("Roll 2 Cloud", "ZONE"),
    "SSLM2": ("Nickel Sim 2", "BLITZ"),
    "UMPA2CONULN": ("Tampa 2 Contain", "ZONE"),
    "SSBUTZO": ("SS Blitz 0", "BLITZ"),
    "UIKEBUTZO": ("Mike Blitz 0", "BLITZ"),
    "8SEAMSHCW2": ("Overload 3 Show 2", "BLITZ"),
    "8SUMSHCW2": ("Overload 3 Show 2", "BLITZ"),
    "COVNGWLLLE": ("Cover 6 Willie", "MATCH"),
    "COVNGWLLLLE": ("Cover 6 Willie", "MATCH"),
    "COVNGWLLLLLE": ("Cover 6 Willie", "MATCH"),
    "CNGWLLLE": ("Cover 6 Willie", "MATCH"),
    "CNGWLLLLE": ("Cover 6 Willie", "MATCH"),
    "CNGWLLLLLE": ("Cover 6 Willie", "MATCH"),
    "IBBUTZG": ("LB Blitz 3", "BLITZ"),
    "IBBIITZG": ("LB Blitz 3", "BLITZ"),
    "UAU8DEEP": ("Man Up 3 Deep", "MAN"),

    # Miami, Philadelphia, SF, Washington additions
    "CAIWARTERS": ("Cover 4 Quarters", "MATCH"),
    "ULKESAMCLOSS8": ("Mike Sam Crash Press", "BLITZ"),
    "CA8BLUZMATCH": ("Cover 3 Buzz Match", "MATCH"),
    "OMJBLECBBUTZ": ("Double CB Blitz", "BLITZ"),
    "DWBLECBBUTZ": ("Double CB Blitz", "BLITZ"),
    "SSBUIZ8": ("SS Blitz 3", "BLITZ"),
    "WSBUTZI": ("WS Blitz 3", "BLITZ"),
    "WSBLITZI": ("WS Blitz 3", "BLITZ"),
    "CA8BLRZ": ("Cover 3 Buzz", "ZONE"),
    "CA8LOCK": ("Cover 3 Lock", "MATCH"),
    "CABWILLIE": ("Cover 6 Willie", "MATCH"),
    "GOHAIFOUT": ("60 Half Out", "BLITZ"),
    "COVELUSHN2": ("Cover 6 Show 2", "MATCH"),
    "SSBUIZO": ("SS Blitz 0", "BLITZ"),
    "21NVERTLURK": ("Cover 2 Invert Lurk", "ZONE"),
    "CA9SAUULE": ("Cover 3 Seam", "MATCH"),
    "ROU28CWJD": ("Roll 2 Cloud", "ZONE"),
    "ROU28CLWJD": ("Roll 2 Cloud", "ZONE"),
    "CAGZONE": ("Cover 6", "MATCH"),
    "CA21NVERT": ("Cover 2 Invert", "ZONE"),
    "CA21MVERT": ("Cover 2 Invert", "ZONE"),
    "IDMJBLETE": ("1 Double TE", "MAN"),
    "IDDJBLETE": ("1 Double TE", "MAN"),
    "IDWBLEWRI": ("1 Double WR1", "MAN"),
    "MLL28CLOUD": ("Roll 2 Cloud", "ZONE"),
    "MUMNC": ("Mug Mac", "BLITZ"),
    "COVNGSW": ("Cover 9 Show 2", "MATCH"),
    "GOUSE": ("60 Base", "BLITZ"),
    "GOWT": ("60 Out", "BLITZ"),
    "USE": ("60 Base", "BLITZ"),
    "8SUMSHOW2": ("Overload 3 Show 2", "BLITZ"),
    "CA3SKY": ("Cover 3 Sky", "ZONE"),
    "COVEUZONE": ("Cover 3 Match", "MATCH"),
    "COVN3C1011D": ("Cover 3 Cloud", "ZONE"),
    "IBBUTZI": ("LB Blitz 1", "BLITZ"),
    "LBB11TZ8": ("LB Blitz 3", "BLITZ"),
    "CA8SEAM": ("Cover 3 Seam", "MATCH"),
    "COFFEEHMSE3SEAM": ("Coffeehouse 3 Seam", "BLITZ"),
    "COVER8CLOLJDSHMN2": ("Cover 3 Cloud Show 2", "ZONE"),
    "SAMCUSII": ("Sam Crash 1", "BLITZ"),
    "SAMCRASH1": ("Sam Crash 1", "BLITZ"),
    "OLBCAGAMS": ("OLB Fire Man", "BLITZ"),
    "OVERLOAD8NICKEL": ("Nickel Overload 3", "BLITZ"),
    "IBBILTZ8": ("LB Blitz 3", "BLITZ"),
    "CA3CIOUDSHCW2": ("Cover 3 Cloud Show 2", "ZONE"),
    "8DOLJBLEWZZ": ("3 Double Buzz", "ZONE"),
    "COVERGIUVNT": ("Cover 6 Invert", "MATCH"),
    "CAIHOLE": ("Cover 1 Hole", "MAN"),
    "FSSUNT8": ("FS Slant 3", "BLITZ"),
    "8DOUBLESKY": ("Cover 3 Sky", "ZONE"),
    "BUTZDTPOP": ("Blitz DT Pop", "BLITZ"),
    "CA3FLRE": ("FS Fire 3", "BLITZ"),
    "CBTEXSLU2": ("CB Tex Sim 2", "BLITZ"),
    "SSMLDBUTZ8": ("SS Middle Blitz 3", "BLITZ"),
    "SSTEXSLU2": ("SS Tex Sim 2", "BLITZ"),
    "CAGSHOW2": ("Cover 9 Show 2", "MATCH"),
    "IDDJBLEWRI": ("1 Double WR1", "MAN"),
    "NTLCOVER8CLMJD": ("Roll Cover 3 Cloud", "ZONE"),
    "SSDDG8": ("SS Dog 3", "BLITZ"),
    "COVN9SAMULE": ("Cover 3 Seam", "MATCH"),
    "COVNAZ0NE": ("Cover 3 Match", "MATCH"),
    "ROLLCOVERSCLU": ("Roll Cover 3 Cloud", "ZONE"),
    "ROU28CLMJD": ("Roll 2 Cloud", "ZONE"),
    "COMAURTERS": ("Cover 4 Quarters", "MATCH"),
    "HOLE": ("Cover 1 Hole", "MAN"),
    "0T4LLSI": ("60 Pinch", "BLITZ"),
    "LBPOP0": ("LB Pop 0", "BLITZ"),
    "UNS": ("Guts", "BLITZ"),
    "GUTS": ("Guts", "BLITZ"),
    "GJIS": ("Guts", "BLITZ"),
    "GOH1FWT": ("60 Half Out", "BLITZ"),
}

OFF_REP_EXTRA = {
    'BOUNCEPAPOSTCURLS': 'BOUNCE PA POST CURLS',
    '14B': 'HB',
    'HBCOUUTERWK': 'HB COUNTER WK',
    'HBCOIJUTERWK': 'HB COUNTER WK',
    'HBSPUTO': 'HB SPLIT O',
    'HBSPLITO': 'HB SPLIT O',
    'STmNGTOSS': 'STRONG TOSS',
    'STETOSS': 'STRONG TOSS',
    'STMGTOSS': 'STRONG TOSS',
    'PWERO': 'POWER O',
    'POWERO': 'POWER O',
    'PAPWERO': 'PA POWER O',
    'PAPOWERO': 'PA POWER O',
    'PASPRINTHBFUT': 'PA SPRINT HB FLAT',
    'PASPRINTHBFLAT': 'PA SPRINT HB FLAT',
    'PASPOT': 'PA SPOT',
    'QBSNEAK': 'QB SNEAK',
    'QBSNEAX': 'QB SNEAK',
    'FBDIVEWEAK': 'FB DIVE WEAK',
    'FBDIVEWEAI': 'FB DIVE WEAK',
    'PAWAGGLE': 'PA WAGGLE',
    'INSIDEZONE': 'INSIDE ZONE',
    'MESHSPOT': 'MESH SPOT',
    'HBPUNCH': 'HB PUNCH',
    'BENCHSWAP': 'BENCH SWAP',
    'READOPTION': 'READ OPTION',
    'SEAMSZ-UUDER': 'SEAMS Z-UNDER',
    'HBPDNER': 'HB POWER',
    'PADAGGER': 'PA DAGGER',
    'MTNRPOBUCKPEEKSLANT': 'MTN RPO BUCK PEEK SLANT',
    'MTNPAFKSCREENSCISSORS': 'MTN PA FK SCREEN SCISSORS',
    'MTNRPOALERTBUBBLE': 'MTN RPO ALERT BUBBLE',
    'FOURVERTICLS': 'FOUR VERTICALS',
    'FOURVERTICALS': 'FOUR VERTICALS',
    'HBSCISSORS': 'HB SCISSORS',
    'DIJO': 'DUO',
    'HBMJO': 'HB DUO',
    'MTNMJO': 'MTN DUO',
    '26DLN': '26 DUO',
    'CLEARMJTZOPTION': 'CLEAROUT OPTION',
    'BOUNCEPABUBBLEFLOOD': 'BOUNCE PA BUBBLE FLOOD',
    'BWUCESHALLM': 'BOUNCE SHALLOW',
    'RETURNSPRINTCOMEUCK': 'RETURN SPRINT COMEBACK',
    'RETURNSPRINTTESCREEN': 'RETURN SPRINT TE SCREEN',
    'SHUFFLEHBOUT': 'SHUFFLE HB OUT',
    'SMASHYDMJBLE': 'SMASH Y DOUBLE',
    'PASWITCHREPLACE': 'PA SWITCH REPLACE',
    'PAHS': 'PA WHEEL',
    'STICK': 'STICK',
    'POSTS': 'POSTS',
    'DWBLESUNT': 'DOUBLES SLANT',
    'RPOALERTMJT': 'RPO ALERT OUT',
    'ZONEFAKEJET': 'ZONE FAKE JET',
    'DEEP CHOICE': 'DEEP CHOICE',
    'HBPWERRT': 'HB POWER',
    'MTNMJTSIDEZOUE': 'MTN OUTSIDE ZONE',
    'YTUIL': 'Y TRAIL',
    'HAMMERDAGGER': 'HAMMER DAGGER',
    'HAILMAW': 'HAIL MARY',
    'HAILMARY': 'HAIL MARY',
    'CNSSDRAG': 'CROSS DRAG',
    'STEPWER': 'SLOT POWER',
    'PATECWSS': 'PA TE CROSS',
    'MTNPWER': 'MTN POWER',
    'YFLATGOAUINE': 'Y FLAT GO LINE',
}

TEAM_MAP = {
    'Arizona': {'off_file': 'Cardinals_playbook-template.xlsx', 'def_file': 'Cardinals_defensive_playbook.xlsx', 'off_team': 'Cardinals', 'def_team': 'ARI'},
    'Atlanta': {'off_file': 'Falcons_playbook-template.xlsx', 'def_file': 'Falcons_defensive_playbook.xlsx', 'off_team': 'Falcons', 'def_team': 'ATL'},
    'Buffalo': {'off_file': 'Bills_playbook-template.xlsx', 'def_file': 'Bills_defensive_playbook.xlsx', 'off_team': 'Bills', 'def_team': 'BUF'},
    'Carolina': {'off_file': 'Panthers_playbook-template.xlsx', 'def_file': 'Panthers_defensive_playbook.xlsx', 'off_team': 'Panthers', 'def_team': 'CAR'},
    'Chicago': {'off_file': 'Bears_playbook-template.xlsx', 'def_file': 'Bears_defensive_playbook.xlsx', 'off_team': 'Bears', 'def_team': 'CHI'},
    'Cincinnati': {'off_file': 'Bengals_playbook-template.xlsx', 'def_file': 'Bengals_defensive_playbook.xlsx', 'off_team': 'Bengals', 'def_team': 'CIN'},
    'Cleveland': {'off_file': 'Browns_playbook-template.xlsx', 'def_file': 'Browns_defensive_playbook.xlsx', 'off_team': 'Browns', 'def_team': 'CLE'},
    'Dallas': {'off_file': 'Cowboys_playbook-template.xlsx', 'def_file': 'Cowboys_defensive_playbook.xlsx', 'off_team': 'Cowboys', 'def_team': 'DAL'},
    'Denver': {'off_file': 'Broncos_playbook-template.xlsx', 'def_file': 'Broncos_defensive_playbook.xlsx', 'off_team': 'Broncos', 'def_team': 'DEN'},
    'Detroit': {'off_file': 'Lions_playbook-template.xlsx', 'def_file': 'Lions_defensive_playbook.xlsx', 'off_team': 'Lions', 'def_team': 'DET'},
    'Green Bay': {'off_file': 'Packers_playbook-template.xlsx', 'def_file': 'Packers_defensive_playbook.xlsx', 'off_team': 'Packers', 'def_team': 'GB'},
    'Indianapolis': {'off_file': 'Colts_playbook-template.xlsx', 'def_file': 'Colts_defensive_playbook.xlsx', 'off_team': 'Colts', 'def_team': 'IND'},
    'Jacksonville': {'off_file': 'Jaguars_playbook-template.xlsx', 'def_file': 'Jaguars_defensive_playbook.xlsx', 'off_team': 'Jaguars', 'def_team': 'JAX'},
    'Kansas City': {'off_file': 'Chiefs_playbook-template.xlsx', 'def_file': 'Chiefs_defensive_playbook.xlsx', 'off_team': 'Chiefs', 'def_team': 'KC'},
    'LAC': {'off_file': 'Chargers_playbook-template.xlsx', 'def_file': 'Chargers_defensive_playbook.xlsx', 'off_team': 'Chargers', 'def_team': 'LAC'},
    'Las Vegas': {'off_file': 'Raiders_playbook-template.xlsx', 'def_file': 'Raiders_defensive_playbook.xlsx', 'off_team': 'Raiders', 'def_team': 'LV'},
    'Miami': {'off_file': 'Dolphins_playbook-template.xlsx', 'def_file': 'Dolphins_defensive_playbook.xlsx', 'off_team': 'Dolphins', 'def_team': 'MIA'},
    'New England': {'off_file': 'Patriots_playbook-template.xlsx', 'def_file': 'Patriots_defensive_playbook.xlsx', 'off_team': 'Patriots', 'def_team': 'NE'},
    'NYG': {'off_file': 'Giants_playbook-template.xlsx', 'def_file': 'Giants_defensive_playbook.xlsx', 'off_team': 'Giants', 'def_team': 'NYG'},
    'NYJ': {'off_file': 'Jets_playbook-template.xlsx', 'def_file': 'Jets_defensive_playbook.xlsx', 'off_team': 'Jets', 'def_team': 'NYJ'},
    'Philadelphia': {'off_file': 'Eagles_playbook-template.xlsx', 'def_file': 'Eagles_defensive_playbook.xlsx', 'off_team': 'Eagles', 'def_team': 'PHI'},
    'San Francisco': {'off_file': '49ers_playbook-template.xlsx', 'def_file': '49ers_defensive_playbook.xlsx', 'off_team': '49ers', 'def_team': 'SF'},
    'Tampa Bay': {'off_file': 'Buccaneers_playbook-template.xlsx', 'def_file': 'Buccaneers_defensive_playbook.xlsx', 'off_team': 'Buccaneers', 'def_team': 'TB'},
    'Baltimore': {'off_file': 'Ravens_playbook-template.xlsx', 'def_file': 'Ravens_defensive_playbook.xlsx', 'off_team': 'Ravens', 'def_team': 'BAL'},
    'Houston': {'off_file': 'Texans_playbook-template.xlsx', 'def_file': 'Texans_defensive_playbook.xlsx', 'off_team': 'Texans', 'def_team': 'HOU'},
    'LAR': {'off_file': 'Rams_playbook-template.xlsx', 'def_file': 'Rams_defensive_playbook.xlsx', 'off_team': 'Rams', 'def_team': 'LAR'},
    'Minnesota': {'off_file': 'Vikings_playbook-template.xlsx', 'def_file': 'Vikings_defensive_playbook.xlsx', 'off_team': 'Vikings', 'def_team': 'MIN'},
    'New Orleans': {'off_file': 'Saints_playbook-template.xlsx', 'def_file': 'Saints_defensive_playbook.xlsx', 'off_team': 'Saints', 'def_team': 'NO'},
    'Pittsburgh': {'off_file': 'Steelers_playbook-template.xlsx', 'def_file': 'Steelers_defensive_playbook.xlsx', 'off_team': 'Steelers', 'def_team': 'PIT'},
    'Seattle': {'off_file': 'Seahawks_playbook-template.xlsx', 'def_file': 'Seahawks_defensive_playbook.xlsx', 'off_team': 'Seahawks', 'def_team': 'SEA'},
    'Tennessee': {'off_file': 'Titans_playbook-template.xlsx', 'def_file': 'Titans_defensive_playbook.xlsx', 'off_team': 'Titans', 'def_team': 'TEN'},
    'Washington': {'off_file': 'Commanders_playbook-template.xlsx', 'def_file': 'Commanders_defensive_playbook.xlsx', 'off_team': 'Commanders', 'def_team': 'WAS'},
}

async def run_ocr_on_image(img, engine, scale=2.5):
    h, w = img.shape[:2]
    up = cv2.resize(img, (0, 0), fx=scale, fy=scale, interpolation=cv2.INTER_LANCZOS4)
    tmp_path = os.path.abspath(f"tmp_ocr_{np.random.randint(100000, 999999)}.png")
    cv2.imwrite(tmp_path, up)
    
    file = await storage.StorageFile.get_file_from_path_async(tmp_path)
    stream = await file.open_async(storage.FileAccessMode.READ)
    decoder = await imaging.BitmapDecoder.create_async(stream)
    bm = await decoder.get_software_bitmap_async()
    res = await engine.recognize_async(bm)
    stream.close()
    try:
        os.remove(tmp_path)
    except Exception:
        pass
        
    words = []
    lines = []
    for line in res.lines:
        lines.append(line.text)
        for word in line.words:
            words.append((word.bounding_rect.y / scale, word.bounding_rect.x / scale, word.bounding_rect.width / scale, word.text))
    return words, lines

def classify_card_badge(im_full, card_idx, is_defense=False):
    h, w = im_full.shape[:2]
    if card_idx == 1:
        x0, x1 = int(w * 0.03), int(w * 0.08)
    elif card_idx == 2:
        x0, x1 = int(w * 0.35), int(w * 0.40)
    else:
        x0, x1 = int(w * 0.68), int(w * 0.73)
    y0, y1 = int(h * 0.72), int(h * 0.82)
    crop = im_full[y0:y1, x0:x1]
    if crop.size == 0:
        return "ZONE" if is_defense else "Pass"
    b, g, r = [float(np.mean(crop[:, :, i])) for i in range(3)]
    
    if is_defense:
        if r > 60 and b > 75 and g < 70 and abs(r - b) < 45:
            return "MATCH"
        if r > 75 and g > 50 and b < 60:
            return "MAN"
        if r > 75 and r > g + 20 and r > b + 25:
            return "BLITZ"
        if b > 80 and b > r + 15:
            return "ZONE"
        return "ZONE"
    else:
        if r > 75 and r > g + 20 and r > b + 20:
            return "Run"
        return "Pass"

def resolve_defense_set(formation, raw_text, fallback_set="Normal"):
    s = str(raw_text or "").upper()
    if formation == "Goal Line":
        if "5-3" in s or "53" in s or "6-3" in s or "63" in s:
            return "5-3"
        return "6-2"
    if formation == "4-3":
        if "EVEN" in s or "6-1" in s or "61" in s: return "Even 6-1"
        if "SOLID" in s: return "Over Solid"
        if "TITE LEO" in s or "TITE" in s: return "Tite Leo"
        if "LEO" in s or "IEO" in s: return "Odd Leo"
        if "UNDER" in s: return "Under"
        if "WIDE" in s: return "Over Wide"
        if "WALK" in s: return "Over Walk"
        return "Over"
    if formation == "Nickel":
        if "WIDE" in s or "2-4 WIDE" in s or "24 WIDE" in s: return "Wide"
        if "2-4 DBL MUG" in s or "24 DBL MUG" in s: return "2-4 Dbl Mug"
        if "2-4 LOAD DBL" in s or "24 LOAD DBL" in s: return "2-4 Load Dbl Mug"
        if "2-4 LOAD MUG" in s or "24 LOAD MUG" in s: return "2-4 Load Mug"
        if "2-4 SINGLE MUG" in s or "24 SINGLE MUG" in s: return "2-4 Single Mug"
        if "2-4 ODD MUG" in s or "24 ODD MUG" in s: return "2-4 Odd Mug"
        if "2-4 CUB" in s or "24 CUB" in s: return "2-4 Cub"
        if "2-4 OVER" in s or "24 OVER" in s: return "2-4 Over"
        if "3-3 DBL MUG" in s or "33 DBL MUG" in s: return "3-3 Dbl Mug"
        if "3-3 LOAD DBL" in s or "33 LOAD DBL" in s: return "3-3 Load Dbl Mug"
        if "3-3 LOAD MUG" in s or "33 LOAD MUG" in s: return "3-3 Load Mug"
        if "3-3 SINGLE MUG" in s or "33 SINGLE MUG" in s: return "3-3 Single Mug"
        if "3-3 ODD" in s or "33 ODD" in s: return "3-3 Odd"
        if "3-3 CUB" in s or "33 CUB" in s: return "3-3 Cub"
        if "3-3 OVER" in s or "33 OVER" in s: return "3-3 Over"
        if "DOUBLE MUG" in s: return "Double Mug"
        if "SINGLE MUG" in s: return "Single Mug"
        if "TRIPLE" in s: return "Triple"
        if "LOAD DBL MUG" in s: return "Load Dbl Mug"
        if "LOAD MUG" in s: return "Load Mug"
        if "LOAD" in s: return "Load"
    if formation == "4-2-5":
        if "WIDE" in s: return "Over Wide"
        if "UNDER" in s: return "Under"
        return "Over G"
    if formation == "2-4-5" or formation == "2-5":
        return "Over Wide"
    if formation == "4-4":
        return "Split"
    if formation == "5-2":
        return "Normal"
    if formation == "Quarter":
        if "OVER" in s or "LOAD" in s: return "Over Load"
        return "Normal"
        
    allowed = catalog_sets.DEFENSE_SETS.get(formation, ["Normal"])
    clean = re.sub(r"[^A-Z0-9]", "", s)
    for fw in ("335", "43", "46", "DIME", "DOLLAR", "GOALLINE", "GOAL", "LINE", "NICKEL", "PREVENT", "BEAR", "34", "QUARTER"):
        clean = clean.replace(fw, "")
        
    if not clean:
        return fallback_set if fallback_set in allowed else allowed[0]
        
    for cand in allowed:
        if re.sub(r"[^A-Z0-9]", "", cand.upper()) == clean:
            return cand
    for cand in allowed:
        c_clean = re.sub(r"[^A-Z0-9]", "", cand.upper())
        if c_clean and c_clean in clean:
            return cand
    cand_map = {re.sub(r"[^A-Z0-9]", "", c.upper()): c for c in allowed}
    m = difflib.get_close_matches(clean, list(cand_map.keys()), n=1, cutoff=0.55)
    if m:
        return cand_map[m[0]]
    return fallback_set if fallback_set in allowed else allowed[0]

def resolve_offense_set(formation, raw_text, fallback_set="Normal"):
    if formation == "H Wing": return "Wing"
    if formation == "Hail Mary": return "Hail Mary"
    if formation == "Goal Line": return "Normal"
    if formation == "Split Backs": return "Offset Wing Close"
    if formation == "Full House": return "Tight"
    if formation == "Split T": return "Tight"
    if formation == "Strong I":
        s_upper = str(raw_text or "").upper()
        if "WING" in s_upper: return "Wing"
        if "CLOSE" in s_upper or "Y OFF" in s_upper: return "Y Off Close"
        return "Pro"
        
    allowed = catalog_sets.OFFENSE_SETS.get(formation, ["Normal"])
    s_upper = str(raw_text or "").upper()
    if "ACE DOUBLE WING" in s_upper or "ACEDOUBLEWING" in s_upper: return "Ace Double Wing"
    if "WING PAIR HEAVY" in s_upper or "WING PAIR" in s_upper: return "Wing Pair Heavy"
    if "WING TIGHT LEFT" in s_upper or "TIGHT LEFT" in s_upper or "WING LEFT" in s_upper: return "Wing Tight Left"
    if "DEUCE CLOSE" in s_upper or "CLOSE DEUCE" in s_upper: return "Deuce Close"
    if "BUNCH X NASTY" in s_upper or "BUNCH X" in s_upper: return "Bunch X Nasty"
    if "TIGHT Y OFF FLEX" in s_upper or "OFF FLEX" in s_upper: return "Tight Y Off Flex"
    if "U OFF TRIPS" in s_upper or ("TRIPS" in s_upper and formation == "Singleback"): return "U Off Trips"
    if "TIGHT DOUBLES" in s_upper or ("DOUBLES" in s_upper and formation == "Singleback"): return "Tight Doubles"
    if "WING CLOSE" in s_upper or "CLOSE WING" in s_upper: return "Wing Close"
    if "U OFF STACK" in s_upper or "STACK" in s_upper: return "U Off Stack"
    if "EAST TIGHT" in s_upper: return "East Tight"
    
    # Gun sets
    if "BUNCH SPREAD" in s_upper: return "Bunch Spread"
    if "BUNCH OPEN TE" in s_upper: return "Bunch Open Te"
    if "WING FLEX" in s_upper: return "Wing Flex Offset Wk"
    if "WING SLOT" in s_upper: return "Wing Slot Offset"
    if "WING TIGHT" in s_upper: return "Wing Tight Nasty"
    if "EMPTY BASE" in s_upper: return "Empty Base Trio"
    if "EMPTY CHIPS NASTY" in s_upper: return "Empty Chips Nasty"
    if "EMPTY CHIPS QUADS" in s_upper: return "Empty Chips Quads"
    if "NORMAL Y OFF CLOSE" in s_upper: return "Normal Y Off Close"
    if "SLOT RIGHT" in s_upper: return "Slot Right"
    if "TIGHT OPEN" in s_upper: return "Tight Open"
    if "TIGHT Y OFF WK" in s_upper: return "Tight Y Off Wk"
    if "TRIPS TE FLEX" in s_upper: return "Trips Te Flex"
    if "TRIPS TE OFFSET" in s_upper: return "Trips Te Offset Wk"
    if "TRIPS HB WK" in s_upper: return "Trips Hb Wk"
    if "Y OFF TRIO" in s_upper: return "Y Off Trio Wk"
    if "Y OFF TRIPS WK" in s_upper: return "Y Off Trips Wk"
    if "Y OFF TRIPS" in s_upper: return "Y Off Trips"
    if "SPREAD" in s_upper: return "Spread"
    if "DOUBLE STACK" in s_upper or "DOUBLESUCK" in s_upper: return "Double Stack"
    if "DOUBLES HB WK" in s_upper: return "Doubles Hb Wk"
    
    clean = re.sub(r"[^A-Z0-9]", "", s_upper)
    for fw in ("SINGLEBACK", "SINGLEBCK", "SINGIEBACK", "SINGLERACK", "GUN", "SHOTGUN", "PISTOL", "IFORM", "WEAK", "FORM", "HAIL", "MARY"):
        clean = clean.replace(fw, "")
        
    if not clean:
        return fallback_set if fallback_set in allowed else allowed[0]
        
    for cand in allowed:
        if re.sub(r"[^A-Z0-9]", "", cand.upper()) == clean:
            return cand
            
    cand_map = {re.sub(r"[^A-Z0-9]", "", c.upper()): c for c in allowed}
    m = difflib.get_close_matches(clean, list(cand_map.keys()), n=1, cutoff=0.50)
    if m:
        return cand_map[m[0]]
    return fallback_set if fallback_set in allowed else allowed[0]

def clean_defense_play(raw_title, badge_type):
    t = unicodedata.normalize('NFKD', str(raw_title or "")).encode('ascii', 'ignore').decode('ascii')
    t = re.sub(r"\s+", " ", t.strip()).upper()
    t = re.sub(r"^[Oo0Ee]\s+", "", t)
    t = re.sub(r"^[-–—]\s*", "", t)
    
    k = re.sub(r"[^A-Z0-9]", "", t)
    if k in DEF_REP_EXTRA:
        return DEF_REP_EXTRA[k]
    if k in catalog_plays.DEF_CANON:
        return catalog_plays.DEF_CANON[k]
    m = difflib.get_close_matches(k, list(catalog_plays.DEF_CANON.keys()), n=1, cutoff=0.7)
    if m:
        return catalog_plays.DEF_CANON[m[0]]
    return (t.title(), badge_type)

def clean_offense_play(raw_title, badge_type):
    t = unicodedata.normalize('NFKD', str(raw_title or "")).encode('ascii', 'ignore').decode('ascii')
    t = re.sub(r"\s+", " ", t.strip()).upper()
    t = re.sub(r"^[Oo0Ee]\s+", "", t)
    t = re.sub(r"^[-–—]\s*", "", t)
    
    k = re.sub(r"[^A-Z0-9]", "", t)
    if k in OFF_REP_EXTRA:
        t = OFF_REP_EXTRA[k]
        k = re.sub(r"[^A-Z0-9]", "", t)
        
    if k in catalog_plays.OFF_CANON:
        return catalog_plays.OFF_CANON[k]
    m = difflib.get_close_matches(k, list(catalog_plays.OFF_CANON.keys()), n=1, cutoff=0.7)
    if m:
        return catalog_plays.OFF_CANON[m[0]]
        
    if "RPO" in t:
        p_type = "RPO"
    elif "PA" in t or "PLAY ACTION" in t:
        p_type = "Play Action"
    elif badge_type == "Run":
        p_type = "Run"
    else:
        p_type = "Pass"
    return (t, p_type)

def detect_formation(text_candidate, is_defense=False, fallback="Gun"):
    if not text_candidate:
        return fallback
    clean = re.sub(r"[^A-Z0-9]", "", str(text_candidate).upper())
    
    if is_defense:
        if any(k in clean for k in ("GOALLINE", "GOAL", "62", "53")):
            return "Goal Line"
        if any(k in clean for k in ("335", "335PENNY", "PENNY")):
            return "3-3-5"
        if any(k in clean for k in ("245", "25", "OVERWIDE")):
            return "2-4-5"
        if "425" in clean:
            return "4-2-5"
        if "44" in clean:
            return "4-4"
        if "52" in clean:
            return "5-2"
        if "43" in clean or "EVEN61" in clean or "61" in clean:
            return "4-3"
        if "34" in clean:
            return "3-4"
        if "46" in clean or "BEAR" in clean:
            return "46"
        if "NICKEL" in clean:
            return "Nickel"
        if "DOLLAR" in clean:
            return "Dollar"
        if "DIME" in clean:
            return "Dime"
        if "PREVENT" in clean or "3DEEP" in clean:
            return "Prevent"
        if "QUARTER" in clean:
            return "Quarter"
        return fallback or "Nickel"
    else:
        if any(k in clean for k in ("GOALLINE", "GOAL")):
            return "Goal Line"
        if any(k in clean for k in ("HAILMARY", "HAIL")):
            return "Hail Mary"
        if "SINGLEBACK" in clean or "SINGIEBACK" in clean or "SINGLEBCK" in clean:
            return "Singleback"
        if "PISTOL" in clean:
            return "Pistol"
        if "GUN" in clean or "SHOTGUN" in clean:
            return "Gun"
        if "HWING" in clean:
            return "H Wing"
        if "IFORM" in clean or "IFORM" in clean:
            return "I Form"
        if "STRONGI" in clean or "STRONG" in clean:
            return "Strong I"
        if "WEAK" in clean:
            return "Weak"
        if "WILDCAT" in clean:
            return "Wildcat"
        if "SPLITBACKS" in clean or "SPLITBACK" in clean:
            return "Split Backs"
        if "SPLITT" in clean:
            return "Split T"
        if "FULLHOUSE" in clean:
            return "Full House"
        return fallback or "Gun"

async def extract_formation_plays(target_dir, team_code, formation_name=None, is_defense=False):
    engine = ocr.OcrEngine.try_create_from_user_profile_languages()
    files = sorted([f for f in os.listdir(target_dir) if f.lower().endswith(('.png', '.jpg'))])
    current_set = "Normal"
    current_formation = formation_name or ("Nickel" if is_defense else "Gun")
    rows = []
    
    for f in files:
        im = cv2.imread(os.path.join(target_dir, f))
        if im is None:
            continue
        h, w = im.shape[:2]
        
        # 1. Pill detection from top bar
        top = im[5:45, :]
        gray_top = cv2.cvtColor(top, cv2.COLOR_BGR2GRAY)
        _, thresh_top = cv2.threshold(gray_top, 120, 255, cv2.THRESH_BINARY)
        cnts, _ = cv2.findContours(thresh_top, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        pills = [cv2.boundingRect(c) for c in cnts if cv2.boundingRect(c)[2] >= 30 and 12 <= cv2.boundingRect(c)[3] <= 30 and cv2.boundingRect(c)[0] < 900]
        if pills:
            pills.sort(key=lambda b: b[2], reverse=True)
            px, py, pw, ph = pills[0]
            crop_pill = top[py:py+ph, px:px+pw]
            padded_pill = cv2.copyMakeBorder(crop_pill, 20, 20, 20, 20, cv2.BORDER_CONSTANT, value=[200,200,200])
            words_pill, lines_pill = await run_ocr_on_image(padded_pill, engine, 2.0)
            if lines_pill and len(lines_pill[0].strip()) >= 3:
                current_set = lines_pill[0].strip()
                if formation_name is None:
                    current_formation = detect_formation(current_set, is_defense=is_defense, fallback=current_formation)
        
        # 2. Bottom pass (bottom-relative to support all screenshot heights/widths)
        bottom = im[-95:, :]
        words, _ = await run_ocr_on_image(bottom, engine, 2.5)
        
        w_scale = w / 972.0
        card_ranges = [
            (int(40 * w_scale), int(340 * w_scale)),
            (int(360 * w_scale), int(670 * w_scale)),
            (int(690 * w_scale), int(990 * w_scale))
        ]
        for c_idx, (x_min, x_max) in enumerate(card_ranges, start=1):
            if c_idx == 1:
                x0_f, x1_f = 0.06, 0.33
            elif c_idx == 2:
                x0_f, x1_f = 0.39, 0.65
            else:
                x0_f, x1_f = 0.72, 0.98
                
            crop_slot = im[int(0.6*h):int(0.95*h), int(x0_f*w):int(x1_f*w)]
            if np.std(crop_slot) < 12.0:
                continue # Blank slot
                
            badge_type = classify_card_badge(im, c_idx, is_defense)
            
            c_words = [t for y, x, width, t in sorted(words) if 54 <= y <= 66 and x_min <= x <= x_max and t not in ("o", "O")]
            raw_title = " ".join(c_words).strip()
            
            sub_words = [t for y, x, width, t in sorted(words) if 67 <= y <= 80 and x_min <= x <= x_max and t not in ("o", "O")]
            sub_text = " ".join(sub_words).strip()
            
            card_crop = im[max(0, h-90):h-5, max(0, x_min-10):min(w, x_max+10)]
            
            is_coach_garbage = lambda s: bool(re.search(r'(CALL|AVG|OC[0-9A-Z]{0,4}LIS|100|AM[OD]S|OA[IM]DS|YDS)', re.sub(r"[^A-Z0-9]", "", unicodedata.normalize('NFKD', str(s or "")).encode('ascii', 'ignore').decode('ascii').upper())))
            
            is_invalid_title = lambda lt: (
                lt in ("PASS", "RUN", "ZONE", "MAN", "BLITZ", "MATCH", "WEAK CLOSE", "BEAR UNDER", "I FORM", "I FORM WING", "I WING", "6-2", "5-3") or
                is_coach_garbage(lt) or
                len(lt) < 3 or
                (sub_text and re.sub(r"[^A-Z0-9]", "", lt) == re.sub(r"[^A-Z0-9]", "", sub_text.upper())) or
                any(x in re.sub(r"[^A-Z0-9]", "", lt) for x in (
                    "GOAL62", "GOALLINE62", "LINE62", "LINE53", "GOAL53",
                    "IINE62", "GOAI62", "GOU62", "GOUIINE62", "GONIINE62",
                    "IINE", "UNE62", "GOUINE", "GOAIINE", "IINF", "IINF62",
                    "GOU6", "GOAI6", "GOUIINF", "62", "53"
                ))
            )
            
            needs_recovery = (not raw_title or len(raw_title) < 2 or is_invalid_title(raw_title))
            
            if needs_recovery:
                # Pass 1: raw 1.5x on card
                _, lines_15 = await run_ocr_on_image(card_crop, engine, 1.5)
                for l in lines_15:
                    lt = re.sub(r"\s+", " ", l.strip()).upper()
                    if not is_invalid_title(lt):
                        raw_title = lt
                        break
                        
                # Pass 2: wiped coach suggestion at 1.5x (wipe top 44px)
                if not raw_title or len(raw_title) < 2 or is_invalid_title(raw_title):
                    wiped = card_crop.copy()
                    wiped[:44, :] = 0
                    _, lines_wiped = await run_ocr_on_image(wiped, engine, 1.5)
                    for l in lines_wiped:
                        lt = re.sub(r"\s+", " ", l.strip()).upper()
                        if not is_invalid_title(lt):
                            raw_title = lt
                            break
                            
                # Pass 3: wiped coach suggestion + threshold 140 at 2.0x
                if not raw_title or len(raw_title) < 2 or is_invalid_title(raw_title):
                    gray = cv2.cvtColor(card_crop, cv2.COLOR_BGR2GRAY)
                    _, th = cv2.threshold(gray, 140, 255, cv2.THRESH_BINARY)
                    th[:44, :] = 0
                    _, lines_th = await run_ocr_on_image(th, engine, 2.0)
                    for l in lines_th:
                        lt = re.sub(r"\s+", " ", l.strip()).upper()
                        if not is_invalid_title(lt):
                            raw_title = lt
                            break
                            
                # Pass 4: lower card band at 1.5x (captures plays like 60 OUT, DIVE, etc.)
                if not raw_title or len(raw_title) < 2 or is_invalid_title(raw_title):
                    lower_band = im[max(0, h-65):h, max(0, x_min-10):min(w, x_max+10)]
                    words_l, lines_l = await run_ocr_on_image(lower_band, engine, 1.5)
                    cand_words = [t for y, x, width, t in sorted(words_l) if 20 <= y <= 45 and t not in ("o", "O", "e", "E") and not any(x in re.sub(r"[^A-Z0-9]", "", t.upper()) for x in ("62", "53", "LINE", "IINE", "GOAL", "GOAI", "GOU"))]
                    cand_str = " ".join(cand_words).strip().upper()
                    if cand_str and not is_invalid_title(cand_str):
                        raw_title = cand_str
                        
                # Pass 5: anchor recovery for short title words (e.g. SPOT, DUO, OUT)
                if not raw_title or len(raw_title) < 2 or is_invalid_title(raw_title):
                    title_box = im[max(0, h-40):max(0, h-18), max(0, x_min+24):min(w, x_min+80)]
                    if title_box.size > 0:
                        pad_anchor = cv2.copyMakeBorder(title_box, 40, 40, 40, 100, cv2.BORDER_CONSTANT, value=[0, 0, 0])
                        cv2.putText(pad_anchor, 'PLAY', (title_box.shape[1] + 25, 52), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 255), 1)
                        _, lines_anchor = await run_ocr_on_image(pad_anchor, engine, 2.5)
                        for la in lines_anchor:
                            cleaned_la = re.sub(r'\s*PLAY\s*$', '', la, flags=re.IGNORECASE).strip().upper()
                            if cleaned_la and not is_invalid_title(cleaned_la):
                                raw_title = cleaned_la
                                break
                                
                # Pass 6: inverted threshold on title crop (recovers high-contrast plays like GUTS, 60 OUT, 60 HALF OUT)
                if not raw_title or len(raw_title) < 2 or is_invalid_title(raw_title):
                    c_title = im[max(0, h-42):max(0, h-15), max(0, x_min):min(w, x_min+140)]
                    if c_title.size > 0:
                        gray_c = cv2.cvtColor(c_title, cv2.COLOR_BGR2GRAY)
                        for th_v in (140, 150):
                            _, th_c = cv2.threshold(gray_c, th_v, 255, cv2.THRESH_BINARY)
                            inv_c = 255 - th_c
                            pad_inv = cv2.copyMakeBorder(inv_c, 30, 30, 30, 80, cv2.BORDER_CONSTANT, value=255)
                            cv2.putText(pad_inv, 'PLAY', (c_title.shape[1] + 15, 45), cv2.FONT_HERSHEY_SIMPLEX, 0.45, 0, 1)
                            _, lines_inv = await run_ocr_on_image(pad_inv, engine, 1.5)
                            for li in lines_inv:
                                cleaned_li = re.sub(r'\s*PLAY\s*$', '', li, flags=re.IGNORECASE).strip().upper()
                                cleaned_li = re.sub(r'^PLAY\s*', '', cleaned_li).strip()
                                if cleaned_li and not is_invalid_title(cleaned_li) and len(cleaned_li) >= 3:
                                    raw_title = cleaned_li
                                    break
                            if raw_title and not is_invalid_title(raw_title):
                                break
                                
                # Pass 7: binary threshold on title crop (recovers high-contrast compact cards like GUTS and 60 HALF OUT)
                if not raw_title or len(raw_title) < 2 or is_invalid_title(raw_title):
                    title_box7 = im[max(0, h-42):max(0, h-18), max(0, x_min+20):min(w, x_min+100)]
                    if title_box7.size > 0:
                        gray7 = cv2.cvtColor(title_box7, cv2.COLOR_BGR2GRAY)
                        _, th7 = cv2.threshold(gray7, 140, 255, cv2.THRESH_BINARY)
                        pad7 = cv2.copyMakeBorder(th7, 30, 30, 30, 80, cv2.BORDER_CONSTANT, value=0)
                        cv2.putText(pad7, 'PLAY', (title_box7.shape[1] + 15, 45), cv2.FONT_HERSHEY_SIMPLEX, 0.45, 255, 1)
                        _, lines7 = await run_ocr_on_image(pad7, engine, 2.0)
                        for l7 in lines7:
                            cl7 = re.sub(r'\s*PLAY\s*$', '', l7, flags=re.IGNORECASE).strip().upper()
                            cl7 = re.sub(r'^PLAY\s*', '', cl7).strip()
                            if cl7 and not is_invalid_title(cl7) and len(cl7) >= 3:
                                raw_title = cl7
                                break
            
            if not raw_title or len(raw_title) < 2 or is_invalid_title(raw_title):
                continue
                
            # Determine card formation
            if formation_name:
                active_formation = formation_name
            else:
                active_formation = detect_formation(sub_text or current_set, is_defense=is_defense, fallback=current_formation)
                current_formation = active_formation
                
            if active_formation == "I Form" and raw_title in ("I FORM", "I FORM WING", "I WING", "IFORMWING", "I Y OFF CLOSE", "I CLOSE", "I OFF CLOSE"):
                continue
            if active_formation == "Goal Line" and any(x in re.sub(r"[^A-Z0-9]", "", raw_title.upper()) for x in ("62", "GOAL62", "GOALLINE62", "LINE62", "IINE62", "GOAI62", "GOU62", "GOUIINE62", "GONIINE62")):
                continue
            if active_formation in ("3-3-5", "335") and re.sub(r"[^A-Z0-9]", "", raw_title.upper()) in ("985", "335", "335PENNY"):
                continue
                
            set_candidate = sub_text if sub_text and len(sub_text) >= 3 else current_set
            
            if is_defense:
                card_set = resolve_defense_set(active_formation, set_candidate, fallback_set=current_set)
                play_name, play_type = clean_defense_play(raw_title, badge_type)
            else:
                card_set = resolve_offense_set(active_formation, set_candidate, fallback_set=current_set)
                play_name, play_type = clean_offense_play(raw_title, badge_type)
                
            rows.append(
                PlayRow(
                    team=team_code,
                    formation=active_formation,
                    set_name=card_set,
                    play_name=play_name,
                    play_type=play_type,
                    source_file=f,
                    card_index=c_idx,
                )
            )
    return rows

def dedupe(rows):
    seen = set()
    unique = []
    for r in rows:
        key = (r.team.upper(), r.formation.upper(), r.set_name.upper(), r.play_name.upper())
        if key not in seen:
            seen.add(key)
            unique.append(r)
    return unique

async def process_team(team_name, out_dir):
    cfg = TEAM_MAP[team_name]
    base_team = os.path.join(r"D:\Robert\Desktop\App Project Builds\Gridiron_OC\Gridiron_OC\imports\Madden 27 Team Playbooks (Screenshots)", team_name)
    
    print(f"\n========================================================")
    print(f"=== Processing Team: {team_name} ===")
    print(f"========================================================")
    
    # Helper to resolve folder case-insensitively
    def find_sub_dir(base, candidates):
        for c in candidates:
            p = os.path.join(base, c)
            if os.path.exists(p):
                return p
        return None

    # 1. Defense
    def_dir = find_sub_dir(base_team, ["Defense", "defense"])
    def_rows = []
    if def_dir and os.path.exists(def_dir):
        direct_images = [f for f in os.listdir(def_dir) if f.lower().endswith(('.png', '.jpg')) and os.path.isfile(os.path.join(def_dir, f))]
        if direct_images:
            print(f"[{team_name}] Scanning Defense (Flat Mode, {len(direct_images)} screenshots found)...")
            f_rows = await extract_formation_plays(def_dir, cfg['def_team'], formation_name=None, is_defense=True)
            print(f"   -> Extracted {len(f_rows)} plays")
            def_rows.extend(f_rows)
        else:
            for form in sorted(os.listdir(def_dir)):
                fp = os.path.join(def_dir, form)
                if os.path.isdir(fp):
                    clean_form = "Goal Line" if form.lower().replace(" ", "") == "goalline" else form
                    print(f"[{team_name}] Scanning Defense formation: {clean_form}...")
                    f_rows = await extract_formation_plays(fp, cfg['def_team'], clean_form, is_defense=True)
                    print(f"   -> Extracted {len(f_rows)} plays")
                    def_rows.extend(f_rows)
                
    unique_def = dedupe(def_rows)
    print(f"[{team_name}] Defense: {len(def_rows)} raw slots -> {len(unique_def)} unique plays")
    
    wb_def = Workbook()
    ws_def = wb_def.active
    ws_def.title = "Defensive Plays"
    ws_def.append(["Team", "Formation", "Set", "Play Name", "Type"])
    for cell in ws_def[1]: cell.font = Font(bold=True)
    for r in sorted(unique_def, key=lambda x: (x.formation, x.set_name, x.play_name)):
        ws_def.append([r.team, r.formation, r.set_name, r.play_name, r.play_type])
        
    summary = wb_def.create_sheet("By Formation Set")
    summary.append(["Formation", "Set", "Play Count"])
    for cell in summary[1]: cell.font = Font(bold=True)
    counts = Counter((r.formation, r.set_name) for r in unique_def)
    for (formation, set_name), count in sorted(counts.items()):
        summary.append([formation, set_name, count])
        
    types = wb_def.create_sheet("By Type")
    types.append(["Type", "Count"])
    for cell in types[1]: cell.font = Font(bold=True)
    type_counts = Counter(r.play_type for r in unique_def)
    for p_type, count in sorted(type_counts.items()):
        types.append([p_type, count])
        
    def_path = os.path.join(out_dir, cfg['def_file'])
    wb_def.save(def_path)
    print(f"Saved: {def_path}")
    
    # 2. Offense
    off_dir = find_sub_dir(base_team, ["Offense", "offense"])
    off_rows = []
    if off_dir and os.path.exists(off_dir):
        direct_images = [f for f in os.listdir(off_dir) if f.lower().endswith(('.png', '.jpg')) and os.path.isfile(os.path.join(off_dir, f))]
        if direct_images:
            print(f"[{team_name}] Scanning Offense (Flat Mode, {len(direct_images)} screenshots found)...")
            f_rows = await extract_formation_plays(off_dir, cfg['off_team'], formation_name=None, is_defense=False)
            print(f"   -> Extracted {len(f_rows)} plays")
            off_rows.extend(f_rows)
        else:
            for form in sorted(os.listdir(off_dir)):
                fp = os.path.join(off_dir, form)
                if os.path.isdir(fp):
                    clean_form = "Goal Line" if form.lower().replace(" ", "") == "goalline" else form
                    print(f"[{team_name}] Scanning Offense formation: {clean_form}...")
                    f_rows = await extract_formation_plays(fp, cfg['off_team'], clean_form, is_defense=False)
                    print(f"   -> Extracted {len(f_rows)} plays")
                    off_rows.extend(f_rows)
                
    unique_off = dedupe(off_rows)
    print(f"[{team_name}] Offense: {len(off_rows)} raw slots -> {len(unique_off)} unique plays")
    
    wb_off = Workbook()
    ws_off = wb_off.active
    ws_off.title = "Playbook"
    ws_off.append(["Team", "Formation", "Set", "Play Name", "Type"])
    for cell in ws_off[1]: cell.font = Font(bold=True)
    for r in sorted(unique_off, key=lambda x: (x.formation, x.set_name, x.play_name)):
        ws_off.append([r.team, r.formation, r.set_name, r.play_name, r.play_type])
        
    off_path = os.path.join(out_dir, cfg['off_file'])
    wb_off.save(off_path)
    print(f"Saved: {off_path}")

async def run_remaining_teams(target_teams=None):
    out_dir = r"D:\Robert\Desktop\App Project Builds\Gridiron_OC\Gridiron_OC\imports\Madden 27 Team Playbook (Excel Workbooks)"
    os.makedirs(out_dir, exist_ok=True)
    teams = target_teams or ['Baltimore', 'Houston', 'LAR', 'Minnesota', 'New Orleans', 'Pittsburgh', 'Seattle', 'Tennessee']
    for t in teams:
        await process_team(t, out_dir)
    print("\n========================================================")
    print("=== ALL TARGET TEAM PLAYBOOKS COMPLETED SUCCESSFULLY! ===")
    print("========================================================")

if __name__ == "__main__":
    target = sys.argv[1:] if len(sys.argv) > 1 else None
    asyncio.run(run_remaining_teams(target))


